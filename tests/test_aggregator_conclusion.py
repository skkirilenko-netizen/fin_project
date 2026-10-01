"""Базовое заключение по данным агрегатора: только агрегатор, состав полный, число с кодом."""

from datetime import date

import pytest
from docx import Document

from finlib.metrics.display import DIGIT_SPACE
from finlib.report.aggregator import (
    BaseConclusion,
    Line,
    NotAggregatorOnlyError,
    Part,
    build,
    render,
)
from finlib.report.policy import AggregatorConclusion
from finlib.scoring.routing import Verdict
from finlib.scoring.routing_store import RoutingRow

SECTIONS = [
    {"code": code, "title": code}
    for code in ("route", "values", "trend", "refinancing", "changes", "limits")
]


WORDING = {
    "basket_senior": "Корзина: {basket}. Старшее основание — {senior}; также: {others}.",
    "basket_single": "Корзина: {basket}. Основание — {senior}.",
    "basket_plain": "Корзина: {basket}.",
    "reference_feature": "{name}: {outcome} — справочно, в маршрут не входит.",
    "reference_outcomes": {"fired": "сработал", "quiet": "не сработал"},
    "rating_after_settlement": (
        "Дефолт по выпуску {issue} урегулирован {settled_on}; рейтинг {point} "
        "({agency}) присвоен {rated_on} — после урегулирования не пересматривался."
    ),
    "rating_revised_after_settlement": (
        "Дефолт по выпуску {issue} урегулирован {settled_on}; рейтинг {point} "
        "({agency}) присвоен {rated_on}, после урегулирования."
    ),
    "bound_meaningless": "долговая нагрузка не определена: прибыль от продаж близка к нулю",
    "ltm_missing": "не сложился: {reason}",
}


def _composition(sections: list[dict] = SECTIONS) -> AggregatorConclusion:
    """Состав из теста, а не из справочника: состав справочника на согласовании."""
    return AggregatorConclusion.model_validate(
        {
            "title": "Базовое заключение",
            "no_class": "Класс не присваивается.",
            "source": "Величины агрегатора.",
            "sections": sections,
            "limitations": ["Примечания не рассматривались."],
            "print_unit": {"okei": "385", "digits": 1},
            "wording": WORDING,
            "origin": "тест",
        }
    )


def test_a_missing_section_is_refused() -> None:
    """Раздел, не объявленный составом, исчез бы молча."""
    with pytest.raises(ValueError, match="разделы"):
        _composition(SECTIONS[:-1])


def test_a_row_with_document_values_is_not_level_one() -> None:
    """Строка, в величинах которой есть документ, базовым заключением не описывается."""
    row = RoutingRow(
        inn="9703024202",
        name="Сегежа",
        report_date=date(2026, 6, 30),
        verdict=Verdict("review", "Разбор", (), "approved"),
        computed=(),
        sources=("PDF", "Cbonds"),
    )
    with pytest.raises(NotAggregatorOnlyError):
        build(row, None, _composition(), date(2026, 9, 29), "Разбор")


def test_every_number_is_printed_with_its_code(tmp_path) -> None:  # noqa: ANN001
    """Строка с величиной печатается таблицей из трёх граф: наименование, значение, код."""
    said = BaseConclusion(
        "7840346335",
        "Группа Илим",
        date(2026, 9, 29),
        date(2025, 12, 31),
        ["Величины агрегатора."],
        [
            Part(
                "values",
                "Величины",
                lines=[Line("Чистый долг / EBITDA", "6,07", "net_debt_ebitda")],
            )
        ],
    )
    path = render(said, _composition(), tmp_path / "проба.docx", "ПРОБА")
    table = Document(str(path)).tables[0]
    assert [cell.text for cell in table.rows[1].cells] == [
        "Чистый долг / EBITDA",
        "6,07",
        "net_debt_ebitda",
    ]


def test_the_reconciliation_share_comes_from_the_table(db_conn) -> None:  # noqa: ANN001
    """Доля совпавших берётся из записанной сверки, а без сверки сборка отказывается.

    Решение владельца 29.09.2026: «176 из 183», вписанное строкой в методику,
    пережило бы перемер сверки.
    """
    from finlib.db import execute
    from finlib.report.aggregator import ReconciliationMissingError, _source

    composition = _composition().model_copy(
        update={"source": "На {issuers} эмитентах совпало {matched} из {compared}."}
    )
    execute("DELETE FROM source_reconciliation", conn=db_conn)
    with pytest.raises(ReconciliationMissingError):
        _source(composition, db_conn)
    for code, outcome, role in (
        ("ifrs.revenue", "совпало", "reporting"),
        ("ifrs.cash", "расходится", "reporting"),
        ("ifrs.ppe", "нет в документе", "reporting"),
        ("ifrs.inventories", "совпало", "comparative"),
    ):
        execute(
            "INSERT INTO source_reconciliation (inn, document_path, report_date, "
            "period_role, line_code, aggregator_kind, outcome, code_version) "
            "VALUES ('7838360491', 'x.pdf', %(d)s, %(r)s, %(c)s, 'exact', %(o)s, 'тест')",
            {"d": date(2025, 12, 31), "r": role, "c": code, "o": outcome},
            conn=db_conn,
        )
    # Сравнительная колонка и «нет в документе» в меру не идут.
    assert _source(composition, db_conn) == "На 1 эмитентах совпало 1 из 2."


def _row(verdict, **kwargs) -> RoutingRow:  # noqa: ANN001
    """Строка маршрута МСФО в тысячах рублей."""
    from finlib.standards import Standard

    return RoutingRow(
        inn="7722514880",
        name="Росинтер",
        report_date=date(2025, 12, 31),
        verdict=verdict,
        computed=(),
        sources=("Cbonds",),
        standard=Standard.IFRS,
        unit="тыс. руб.",
        unit_code="384",
        **kwargs,
    )


def test_basket_names_the_senior_ground_and_the_rest() -> None:
    """«Разбор (рынок)» молчал о величинах; теперь — старшее и остальные."""
    from finlib.report.aggregator import _route

    verdict = Verdict(
        "review", "Разбор", (), "approved",
        subgroup_names=("рынок", "риск по величинам"),
    )
    part = Part("route", "Вывод")
    _route(part, _row(verdict), None, _composition(), date(2026, 10, 1), "Разбор")
    assert part.paragraphs[0] == (
        "Корзина: Разбор. Старшее основание — рынок; также: риск по величинам."
    )


def test_ground_money_is_printed_in_the_document_unit() -> None:
    """Платежи 127 760 и деньги 1 728 тыс. руб. в основании — 127,8 и 1,7 млн руб."""
    from decimal import Decimal

    from finlib.report.aggregator import _route
    from finlib.scoring.routing import Refinance, route

    verdict = route(
        (),
        unit="тыс. руб.",
        quarantined=False,
        today=date(2026, 10, 1),
        latest_annual=date(2025, 12, 31),
        refinance=Refinance(Decimal("127760"), Decimal("1728"), "тыс. руб.", 365),
    )
    gap = next(item for item in verdict.findings if item.ground == "refinancing_gap")
    # Список наблюдения печатает как прежде — в единице комплекта.
    assert f"127{DIGIT_SPACE}760" in gap.text and "тыс. руб." in gap.text
    part = Part("route", "Вывод")
    _route(part, _row(verdict), None, _composition(), date(2026, 10, 1), "Разбор")
    said = next(text for text in part.paragraphs if "refinancing_gap" in text)
    assert "127,8" in said and "1,7 млн руб." in said
    assert "тыс. руб." not in said


def test_values_money_in_document_unit_and_ratio_as_is() -> None:
    """Чистый долг 2 469 795 тыс. руб. — 2 469,8 млн руб.; отношение не переводится."""
    from decimal import Decimal

    from finlib.report.aggregator import _values

    row = _row(
        Verdict("review", "Разбор", (), "approved"),
        shown_values=(
            ("net_debt", "Чистый долг", "2 469 795 тыс. руб."),
            ("net_debt_ebitda", "Чистый долг / EBITDA", "1,34"),
        ),
        values={"net_debt": Decimal("2469795"), "net_debt_ebitda": Decimal("1.34")},
    )
    part = Part("values", "Величины")
    _values(part, row, None, _composition(), date(2026, 10, 1), "Разбор")
    assert [line.shown for line in part.lines] == [f"2{DIGIT_SPACE}469,8 млн руб.", "1,34"]
