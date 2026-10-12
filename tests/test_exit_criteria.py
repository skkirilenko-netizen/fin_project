"""Критерий выхода «эмитенты закрыты»: пять критериев на синтетике в findb_test."""

import json
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from finlib.config import settings
from finlib.db import PgConnection, execute, fetch_all

sys.path.insert(0, str(settings.base_dir / "eval"))
import exit_criteria as exits  # noqa: E402

DAY = date(2090, 3, 15)  # среда
INNS = [f"00000000{number:02d}" for number in range(80, 86)]
MARKS = "вердикт эксперта (согласен / не согласен / лишний / нет выпусков)"
DEFECT = "критичный дефект (да / нет)"


def _point(conn: PgConnection, inn: str, day: date, basket: str, kind: str = "run") -> None:
    """Точка истории маршрута; у наблюдения — свой прогон дня."""
    run_id = None
    if kind == "run":
        [row] = fetch_all(
            "INSERT INTO routing_run (kind, as_of, status) VALUES ('run', %(day)s, 'done') "
            "RETURNING id",
            {"day": day},
            conn=conn,
        )
        run_id = row["id"]
    execute(
        "INSERT INTO routing_history (run_id, inn, as_of, kind, basket, fingerprint) "
        "VALUES (%(run)s, %(inn)s, %(day)s, %(kind)s, %(basket)s, 'fp')",
        {"run": run_id, "inn": inn, "day": day, "kind": kind, "basket": basket},
        conn=conn,
    )


@pytest.fixture
def conn(db_conn: PgConnection) -> PgConnection:
    """Чистая история синтетических ИНН и прогонов 2090 года в откатываемой транзакции."""
    execute("DELETE FROM routing_history WHERE inn = ANY(%(inns)s)", {"inns": INNS},
            conn=db_conn)
    execute("DELETE FROM routing_history WHERE run_id IN "
            "(SELECT id FROM routing_run WHERE as_of >= '2089-01-01')", {}, conn=db_conn)
    execute("DELETE FROM routing_run WHERE as_of >= '2089-01-01'", {}, conn=db_conn)
    return db_conn


def test_months_back_keeps_the_day_or_takes_the_month_end() -> None:
    """Месяц назад от 31 марта — последний день февраля; год назад — тот же день."""
    assert exits.months_back(date(2090, 3, 31), 1) == date(2090, 2, 28)
    assert exits.months_back(date(2090, 1, 15), 1) == date(2089, 12, 15)
    assert exits.months_back(date(2090, 3, 15), 12) == date(2089, 3, 15)


def test_coverage_names_who_has_no_basket(conn: PgConnection) -> None:
    """Охват: с корзиной, вышедший записью и без корзины — последний назван и рушит критерий."""
    _point(conn, INNS[0], DAY, "clear")
    _point(conn, INNS[1], DAY, "status_unknown")
    universe = {inn: inn for inn in INNS[:4]}
    found = exits.coverage(conn, DAY, universe, {INNS[2]})
    assert found.done is False
    assert "с корзиной на 15.03.2090 2 (из них «Установить статус» 1)" in found.fact
    assert "вышли записью в журнале исключений 1" in found.fact
    assert found.fact.endswith(f"без корзины 1: {INNS[3]}")
    whole = exits.coverage(conn, DAY, {inn: inn for inn in INNS[:3]}, {INNS[2]})
    assert whole.done is True


def test_coverage_without_the_list_or_the_route_is_not_established(conn: PgConnection) -> None:
    """Перечня на диске нет либо маршрута дня нет — «не установлено», а не «нет»."""
    assert exits.coverage(conn, DAY, None, set()).done is None
    assert exits.coverage(conn, DAY, {INNS[0]: ""}, set()).done is None


def test_quality_catches_clear_on_the_eve_by_both_kinds(conn: PgConnection) -> None:
    """Качество: «Без внимания» накануне события — по пересчёту и по наблюдению отдельно."""
    event = DAY - timedelta(days=10)
    _point(conn, INNS[0], event - timedelta(days=7), "clear", kind="backfill")
    _point(conn, INNS[1], event - timedelta(days=1), "attention")
    _point(conn, INNS[1], event, "clear")  # день события — уже не накануне
    happened = {
        INNS[0]: event,
        INNS[1]: event,
        INNS[2]: DAY - timedelta(days=400),  # вне двенадцати месяцев
    }
    found = exits.quality(conn, DAY, happened)
    assert found.done is False
    assert "событий с 16.03.2089 по 15.03.2090: 2" in found.fact
    assert f"«Без внимания» накануне 1 — {INNS[0]} (05.03.2090)" in found.fact
    assert "наблюдение: с историей до события 1, «Без внимания» накануне 0" in found.fact


def test_quality_without_history_is_not_established(conn: PgConnection) -> None:
    """Событие есть, а истории до него нет — «не установлено»; календаря нет — тоже."""
    assert exits.quality(conn, DAY, {INNS[0]: DAY}).done is None
    assert exits.quality(conn, DAY, None).done is None


def _scheduled_runs(conn: PgConnection, days: list[date], failed: date | None = None) -> None:
    """Прогоны по расписанию (10:00:30 по местному времени) с доставками."""
    for day in days:
        started = datetime(day.year, day.month, day.day, 10, 0, 30).astimezone()
        status = "failed" if day == failed else "done"
        execute(
            "INSERT INTO routing_run (kind, as_of, status, started_at, sources) "
            "VALUES ('run', %(day)s, 'done', %(started)s, %(sources)s)",
            {"day": day, "started": started,
             "sources": json.dumps([{"name": "рейтинги", "status": status}])},
            conn=conn,
        )


def _weekdays(since: date, until: date) -> list[date]:
    """Рабочие дни (пн–пт) от since до until включительно."""
    days = [since + timedelta(days=offset) for offset in range((until - since).days + 1)]
    return [day for day in days if day.weekday() < 5]


def test_stability_needs_a_month_of_clean_scheduled_runs(conn: PgConnection) -> None:
    """Месяц чистых прогонов по расписанию — выполнен; ручной прогон в счёт не идёт."""
    month = _weekdays(date(2090, 2, 16), DAY)
    _scheduled_runs(conn, month)
    found = exits.stability(conn, DAY)
    assert found.done is True
    assert f"{len(month)} рабочих дней (нужно {len(month)}" in found.fact


def test_stability_is_broken_by_a_failed_delivery(conn: PgConnection) -> None:
    """Отказ доставки в прогоне по расписанию обрывает счёт."""
    month = _weekdays(date(2090, 2, 16), DAY)
    _scheduled_runs(conn, month, failed=date(2090, 3, 9))
    found = exits.stability(conn, DAY)
    assert found.done is False
    assert "09.03.2090: прогон по расписанию не чистый" in found.fact


def _answer(
    path: Path,
    day: date,
    review: list[tuple[str, str]],
    sample: list[tuple[str, str]] | None = None,
    clear: list[tuple[str, str]] | None = None,
    seen: str = "да",
) -> Path:
    """Файл ответа эксперта вида пакета v2: листы, графы вердикта, отметка."""
    book = Workbook()
    book.remove(book.active)
    instruction = book.create_sheet("Инструкция")
    instruction.append(["Отметка эксперта", None])
    instruction.append([exits.SIGN_CLEAR_SEEN, seen])
    book.create_sheet("Сводка").append(["Дата маршрута", f"{day:%d.%m.%Y}"])
    for title, rows in (
        ("Разбор", review),
        ("Без внимания", clear if clear is not None else [("", "")]),
        ("Выборка", sample if sample is not None else [("согласен", "")]),
    ):
        sheet = book.create_sheet(title)
        sheet.append(["ИНН", "корзина", MARKS, DEFECT, "комментарий"])
        for number, (mark, defect) in enumerate(rows):
            sheet.append([f"77{number:08d}", title, mark or None, defect or None, None])
    book.save(str(path))
    return path


def test_expertise_needs_two_passed_reviews_in_a_row(tmp_path: Path) -> None:
    """Две последние по дате проверки пройдены — выполнен; более ранняя не в счёт."""
    good = [("согласен", "нет")] * 9 + [("лишний", "")]
    _answer(tmp_path / "a.xlsx", date(2090, 3, 1), [("согласен", "да")])
    _answer(tmp_path / "b.xlsx", date(2090, 3, 8), good)
    _answer(tmp_path / "c.xlsx", date(2090, 3, 15), good)
    found = exits.expertise(tmp_path)
    assert found.done is True
    assert "01.03.2090 (a.xlsx): не пройдена; критичных 1" in found.fact
    assert "«лишних» в «Разборе» 1 из 10" in found.fact


def test_expertise_fails_on_a_critical_defect_or_too_many_extra(tmp_path: Path) -> None:
    """Критичный дефект в «Без внимания» либо «лишних» больше 10 % — проверка не пройдена."""
    _answer(tmp_path / "a.xlsx", date(2090, 3, 8), [("согласен", "")],
            clear=[("не согласен", "да")])
    _answer(tmp_path / "b.xlsx", date(2090, 3, 15), [("согласен", "")] * 4 + [("лишний", "")])
    found = exits.expertise(tmp_path)
    assert found.done is False
    reviews = [exits.read_review(tmp_path / name) for name in ("a.xlsx", "b.xlsx")]
    assert [item.passed for item in reviews] == [False, False]
    assert reviews[0].critical == 1 and reviews[1].extra == 1


def test_an_unmarked_or_unknown_answer_makes_the_review_incomplete(tmp_path: Path) -> None:
    """Пустой вердикт в «Разборе», слово вне перечня и неотмеченный просмотр — не завершена."""
    empty = exits.read_review(_answer(tmp_path / "a.xlsx", DAY, [("", "")]))
    assert not empty.complete and "«Разбор»: не размечено 1 из 1" in empty.why
    odd = exits.read_review(_answer(tmp_path / "b.xlsx", DAY, [("лишнее", "")]))
    assert not odd.complete and "слова вне перечня графы — лишнее" in odd.why
    unseen = exits.read_review(_answer(tmp_path / "c.xlsx", DAY, [("согласен", "")], seen=""))
    assert not unseen.complete and "«Без внимания» просмотрен целиком" in unseen.why
    assert not unseen.passed


def test_expertise_without_answers_is_not_established(tmp_path: Path) -> None:
    """Файлов ответов нет — «не установлено»."""
    assert exits.expertise(tmp_path / "нет").done is None


def test_documentation_names_every_ground_and_is_not_a_draft(tmp_path: Path) -> None:
    """Документ называет все коды и не черновик — выполнен; иначе называет недостающие."""
    routing = tmp_path / "methodology" / "routing.yaml"
    routing.parent.mkdir()
    routing.write_text(
        "baskets:\n  - code: review\n    grounds:\n      - code: emission_default\n"
        "reference:\n  - code: muted_by_branch\n"
        "issuer_types:\n  - code: structural\n    ground: structural_issuer\n"
        "ground_sources:\n  payment_missed: выпуск\n",
        encoding="utf-8",
    )
    doc = tmp_path / "docs" / "overview.md"
    doc.parent.mkdir()
    named = "`emission_default` `muted_by_branch` `structural_issuer` `payment_missed`"
    doc.write_text(f"# Обзор\n\n{named}\n", encoding="utf-8")
    assert exits.documentation(doc, routing).done is True
    doc.write_text(f"# Обзор\n\n> **Черновик**\n\n{named}\n", encoding="utf-8")
    assert "помечен черновиком" in exits.documentation(doc, routing).fact
    doc.write_text("# Обзор\n\n`emission_default`\n", encoding="utf-8")
    found = exits.documentation(doc, routing)
    assert found.done is False and "названо кодом 1 из 4" in found.fact


def test_the_table_prints_every_criterion_and_the_total(conn: PgConnection, tmp_path: Path) -> None:
    """Таблица: пять строк, исход словом, итог «выполнено N из 5»."""
    rows = exits.build(conn, DAY, tmp_path, tmp_path / "docs" / "нет.md", None, set(), None)
    text = exits.render(DAY, rows)
    assert [row.name for row in rows] == [
        "охват", "качество", "устойчивость", "экспертиза", "документация"]
    assert text.count("| не установлено |") == 3
    assert "**Выполнено 0 из 5**" in text


def test_an_existing_result_is_not_overwritten(tmp_path: Path, capsys) -> None:  # noqa: ANN001
    """Готовый файл итога не переписывается: отказ до чтения диска и базы."""
    out = tmp_path / "exit.md"
    out.write_text("x", encoding="utf-8")
    assert exits.main(["--as-of", "2090-03-15", "--out", str(out)]) == 1
    assert "не переписывается" in capsys.readouterr().out
    assert out.read_text(encoding="utf-8") == "x"
