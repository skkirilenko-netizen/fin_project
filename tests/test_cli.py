"""Тесты CLI: вывод по этапам и внятная остановка с указанием причины.

Сеть и модель не задействованы: полный цикл проверяется через pipeline
на подставном источнике, а команды чтения — на загруженных пробах.
"""

import pytest
from typer.testing import CliRunner

from finlib.cli import app
from finlib.pipeline import PipelineError, Stage

FULL_INN = "7736050003"
STOPPED_INN = "2100010824"
NO_CLASS_INN = "2522002003"

runner = CliRunner()


# --- проверка ввода ---------------------------------------------------------


@pytest.mark.parametrize("inn", ["123", "abcdefghij", "77360500031"])
def test_bad_inn_is_rejected_before_any_work(inn: str) -> None:
    """Негодный ИНН отсекается до обращения к сети и базе."""
    result = runner.invoke(app, ["show", "--inn", inn])
    assert result.exit_code == 1
    assert "не похож на настоящий" in result.output


def test_unknown_inn_says_what_to_do(db_conn) -> None:
    """По организации без расчётов команда говорит, что сделать дальше."""
    result = runner.invoke(app, ["show", "--inn", "0000000000"])
    assert result.exit_code == 1
    assert "analyze --inn" in result.output


# --- show -------------------------------------------------------------------


def test_show_prints_metrics_and_verdict(db_conn) -> None:
    """Таблица показателей и вердикт под ней."""
    result = runner.invoke(app, ["show", "--inn", FULL_INN])
    assert result.exit_code == 0
    assert "Коэффициент текущей ликвидности" in result.output
    assert "Класс: C" in result.output
    assert "Уверенность в оценке" in result.output


def test_show_hides_derived_values(db_conn) -> None:
    """Производные величины в сводку не идут: их место в приложении."""
    result = runner.invoke(app, ["show", "--inn", FULL_INN])
    assert "_chg_pct" not in result.output
    assert "_share" not in result.output


def test_show_withholds_score_without_scoring(db_conn) -> None:
    """Класс от стоп-фактора при узком основании балла не раскрывает."""
    result = runner.invoke(app, ["show", "--inn", STOPPED_INN])
    assert result.exit_code == 0
    assert "Класс: E" in result.output
    assert "балл" not in result.output.split("Класс: E")[1].split("\n")[0]
    assert "Балльная оценка не формируется" in result.output


def test_show_reports_absence_of_class(db_conn) -> None:
    """Отсутствие класса — штатный исход, и он назван причиной."""
    result = runner.invoke(app, ["show", "--inn", NO_CLASS_INN])
    assert result.exit_code == 0
    assert "Класс не присвоен" in result.output


# --- quality ----------------------------------------------------------------


def test_quality_lists_checks_with_counts(db_conn) -> None:
    """Журнал контролей выводится со сводкой по блокирующим."""
    result = runner.invoke(app, ["quality", "--inn", FULL_INN])
    assert result.exit_code == 0
    assert "balance_equality" in result.output
    assert "блокирующих" in result.output


def test_quality_failed_only_narrows_output(db_conn) -> None:
    """Флаг --failed-only оставляет только сработавшие контроли."""
    full = runner.invoke(app, ["quality", "--inn", NO_CLASS_INN])
    narrow = runner.invoke(app, ["quality", "--inn", NO_CLASS_INN, "--failed-only"])
    assert narrow.exit_code == 0
    assert len(narrow.output.split("\n")) < len(full.output.split("\n"))
    assert "pass " not in narrow.output


# --- остановка --------------------------------------------------------------


def test_pipeline_error_names_stage_and_reason() -> None:
    """Ошибка цикла называет этап и причину, а не только факт неудачи."""
    error = PipelineError(Stage.QUALITY, "все комплекты отбракованы")
    assert error.stage is Stage.QUALITY
    assert "контроли качества" in str(error)
    assert "все комплекты отбракованы" in str(error)


def test_stage_order_matches_the_pipeline() -> None:
    """Этапы перечислены в порядке выполнения."""
    assert list(Stage) == [
        Stage.FETCH,
        Stage.LOAD,
        Stage.QUALITY,
        Stage.METRICS,
        Stage.SCORING,
        Stage.CONCLUSION,
        Stage.DOCUMENT,
    ]


def test_analyze_reports_the_failing_stage(monkeypatch) -> None:
    """При остановке цикла команда печатает этап и причину."""
    import finlib.cli as module

    def failing(*args, **kwargs):
        raise PipelineError(Stage.FETCH, "организация не найдена")

    monkeypatch.setattr(module, "analyze", failing)
    result = runner.invoke(app, ["analyze", "--inn", FULL_INN])
    assert result.exit_code == 1
    assert "получение отчётности" in result.output
    assert "организация не найдена" in result.output


# --- пересчёт без источника -------------------------------------------------


def test_reprocess_skips_the_source(tmp_path) -> None:
    """Пересчёт идёт по загруженным данным и к источнику не обращается."""
    result = runner.invoke(
        app,
        ["reprocess", "--inn", STOPPED_INN, "--output", str(tmp_path)],
    )
    assert result.exit_code == 0, result.output
    assert "пропущено: пересчёт из ранее загруженных данных" in result.output
    assert "расчёт показателей" in result.output
    assert list(tmp_path.glob("*.docx")), "документ не записан"


def test_document_without_text_says_so(tmp_path) -> None:
    """Документ, собранный расчётом, оговаривает это прямо.

    Сборка расчётом — режим по умолчанию, поэтому флага в вызове нет:
    обращение к модели включается `--llm`.
    """
    from docx import Document

    from finlib.report.document import CALCULATED_TEXT_NOTICE

    result = runner.invoke(
        app,
        ["reprocess", "--inn", STOPPED_INN, "--output", str(tmp_path)],
    )
    assert result.exit_code == 0, result.output
    path = next(iter(tmp_path.glob("*.docx")))
    text = "\n".join(item.text for item in Document(path).paragraphs)
    assert CALCULATED_TEXT_NOTICE in text
    assert "Ключевой вывод" in text, "расчётная часть остаётся полной"


def test_model_is_off_by_default(tmp_path, monkeypatch) -> None:
    """Без `--llm` к модели не обращаются вовсе.

    Умолчание развёрнуто по замеру 17.09.2026: сборка расчётом даёт документ
    всегда и за секунды, а вклад модели укладывается в связки одного раздела.
    Сторож здесь грубый и потому надёжный — обращение к модели поднимает
    исключение, и молчаливого вызова не останется.
    """
    import finlib.llm.client as client_module

    def refuse(*args, **kwargs):
        raise AssertionError("модель вызвана без --llm")

    monkeypatch.setattr(client_module.LLMClient, "complete", refuse)
    result = runner.invoke(
        app, ["reprocess", "--inn", STOPPED_INN, "--output", str(tmp_path)]
    )
    assert result.exit_code == 0, result.output
    assert list(tmp_path.glob("*.docx")), "документ не записан"


def test_document_without_text_records_no_model(tmp_path) -> None:
    """В приложении честно сказано, что модель не привлекалась."""
    from docx import Document

    result = runner.invoke(
        app,
        ["reprocess", "--inn", STOPPED_INN, "--output", str(tmp_path)],
    )
    assert result.exit_code == 0, result.output
    path = next(iter(tmp_path.glob("*.docx")))
    text = "\n".join(item.text for item in Document(path).paragraphs)
    assert "Языковая модель текстовой части: не привлекалась" in text
