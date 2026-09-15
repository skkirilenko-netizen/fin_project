"""Тесты сборки документа: разбор разделов и запись docx без обращения к модели."""

from datetime import date, datetime

import pytest
from docx import Document

from finlib.llm.service import Conclusion, with_corrections
from finlib.report.document import DISCLAIMER, build_report, output_path
from finlib.report.sections import (
    EXPECTED,
    MissingSectionError,
    split_sections,
)

FULL_INN = "7736050003"
SIMPLE_INN = "2100010824"
NO_CLASS_INN = "2522002003"

ANSWER = """### 2. Фактическая база

Валюта баланса (1600) — 25 736 328 136 тыс. руб.

### 3. Аналитическая интерпретация

Коэффициент автономии (equity_ratio) вырос с 0,62 до 0,64.

### 4. Риски

Отрицательный чистый оборотный капитал (nwc).

### 5. Ограничения анализа

- Отраслевой привязки нет.
- Ряд опирается на сравнительные периоды.

### 6. Вопросы к организации

1. Что произошло с дебиторской задолженностью?
"""


# --- разбор ответа модели ---------------------------------------------------


def test_all_sections_are_parsed() -> None:
    """Пять разделов разбираются по заголовкам."""
    sections = split_sections(ANSWER)
    assert [item.number for item in sections] == [number for number, _ in EXPECTED]
    assert sections[0].title == "Фактическая база"


def test_paragraphs_are_kept() -> None:
    """Содержимое раздела не теряется."""
    sections = {item.number: item for item in split_sections(ANSWER)}
    assert "25 736 328 136" in sections[2].paragraphs[0]
    assert len(sections[5].paragraphs) == 2


def test_list_markers_become_dashes() -> None:
    """Маркеры списка приводятся к единому виду."""
    sections = {item.number: item for item in split_sections(ANSWER)}
    assert sections[5].paragraphs[0].startswith("— ")


def test_missing_section_is_an_error() -> None:
    """Пропущенный раздел — ошибка, а не повод отдать неполный документ."""
    without_risks = ANSWER.replace("### 4. Риски", "### 9. Прочее")
    with pytest.raises(MissingSectionError) as info:
        split_sections(without_risks)
    assert 4 in info.value.missing


def test_section_order_is_ours_not_the_model_s() -> None:
    """Порядок разделов задаём мы: перестановка в документе недопустима."""
    shuffled = "\n".join(
        [
            ANSWER[ANSWER.index("### 6.") :],
            ANSWER[ANSWER.index("### 2.") : ANSWER.index("### 6.")],
        ]
    )
    assert [item.number for item in split_sections(shuffled)] == [2, 3, 4, 5, 6]


# --- повторная попытка ------------------------------------------------------


def test_retry_carries_the_reasons() -> None:
    """Повтор при нулевой температуре осмыслен, только если сказать, что не так."""
    asked = with_corrections("инструкция", ["14,7 — отсутствует во входных данных"])
    assert "Предыдущий ответ отклонён" in asked
    assert "14,7" in asked
    assert asked.startswith("инструкция")


def test_retry_without_reasons_changes_nothing() -> None:
    """Первая попытка идёт с исходной инструкцией."""
    assert with_corrections("инструкция", []) == "инструкция"


# --- документ ---------------------------------------------------------------


def test_output_path_follows_the_agreed_shape(tmp_path) -> None:
    """Имя файла — ИНН и отчётная дата."""
    path = output_path(FULL_INN, date(2025, 12, 31), tmp_path)
    assert path.name == "7736050003_2025-12-31.docx"


@pytest.fixture
def rendered(db_conn, tmp_path):
    """Документ, собранный на готовом тексте, без обращения к модели."""
    conclusion = Conclusion(
        inn=FULL_INN,
        report_date=date(2025, 12, 31),
        text=ANSWER,
        model="тестовая-модель",
        attempt=1,
        checked_numbers=3,
    )
    report = build_report(
        FULL_INN,
        db_conn,
        conclusion=conclusion,
        directory=tmp_path,
        generated_at=datetime(2026, 1, 1, 12, 0),
    )
    return report, Document(report.path)


def document_text(document) -> str:
    """Весь текст документа, включая таблицы."""
    parts = [item.text for item in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_document_is_written(rendered) -> None:
    """Файл создан по согласованному пути."""
    report, _ = rendered
    assert report.path.exists()
    assert report.path.suffix == ".docx"


def test_disclaimer_is_in_the_header(rendered) -> None:
    """Дисклеймер об автоматическом формировании стоит в шапке."""
    _, document = rendered
    assert DISCLAIMER in document_text(document)


def test_all_sections_are_in_the_document(rendered) -> None:
    """Раздел 1, разделы модели и приложение — все на месте."""
    _, document = rendered
    text = document_text(document)
    assert "1. Ключевой вывод" in text
    for number, title in EXPECTED:
        assert f"{number}. {title}" in text
    assert "Приложение" in text


def test_appendix_carries_the_tables(rendered) -> None:
    """В приложении есть таблицы показателей и контролей."""
    _, document = rendered
    text = document_text(document)
    assert "Показатели за периоды" in text
    assert "Выполненные контроли качества" in text
    assert len(document.tables) >= 2


def test_model_name_and_versions_are_recorded(rendered) -> None:
    """Документ говорит, какой моделью и по какой методике сделан."""
    _, document = rendered
    text = document_text(document)
    assert "тестовая-модель" in text
    assert "Версия методики оценки" in text
    assert "01.01.2026" in text


def test_model_text_reaches_the_document(rendered) -> None:
    """Текст модели попадает в документ, а не теряется."""
    _, document = rendered
    assert "25 736 328 136" in document_text(document)


def test_document_without_class_renders(db_conn, tmp_path) -> None:
    """Организация без класса рендерится штатно, а не падает."""
    conclusion = Conclusion(
        inn=NO_CLASS_INN,
        report_date=date(2024, 12, 31),
        text=ANSWER,
        model="тестовая-модель",
        attempt=1,
        checked_numbers=0,
    )
    report = build_report(
        NO_CLASS_INN, db_conn, conclusion=conclusion, directory=tmp_path
    )
    text = document_text(Document(report.path))
    assert "Класс финансового состояния не присвоен" in text
    assert "Балл:" not in text
