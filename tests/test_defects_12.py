"""Тесты семи дефектов, найденных экспертной оценкой пилотных заключений.

Каждый тест назван по номеру дефекта из TASKS-2.md, чтобы регресс был виден
сразу: по имени понятно, что именно вернулось.
"""

from datetime import date, datetime
from decimal import Decimal

import pytest

from finlib.metrics.definitions import EXCLUSION_ORDER, ExclusionKind, load_metrics
from finlib.normalize.lines import UnitSource, load_lines
from finlib.quality.codes import CheckCode
from finlib.quality.thresholds import load_thresholds
from finlib.report.appendix import (
    _percent,
    checks_table,
    groups_table,
    metrics_table,
    not_calculated_table,
    provenance,
)
from finlib.report.consistency import check_document
from finlib.report.data import load_report_data
from finlib.scoring.definitions import load_scoring

FULL_INN = "7736050003"
STOPPED_INN = "2100010824"
NO_CLASS_INN = "2522002003"

LINES = load_lines()
METRICS = load_metrics()
SCORING = load_scoring()


# --- 1. Единица измерения ---------------------------------------------------


def test_defect_1_unit_is_determined_by_the_form() -> None:
    """Единица берётся из формы, а не принимается как предположение."""
    units = LINES.units
    assert units.okei_code == "384"
    assert units.source_for(["0710001", "0710002"]) is UnitSource.FORM_STANDARD
    assert units.origin.strip(), "происхождение правила обязано быть названо"


def test_defect_1_unknown_forms_leave_unit_undetermined() -> None:
    """Комплект из неизвестных форм единицы не получает."""
    assert LINES.units.source_for(["9999999"]) is UnitSource.UNKNOWN
    assert LINES.units.source_for([]) is UnitSource.UNKNOWN


def test_defect_1_control_is_blocking() -> None:
    """Неопределённая единица отправляет комплект в карантин."""
    severity = load_thresholds().severity_of(CheckCode.UNIT_NOT_DETERMINED.value)
    assert severity.value == "blocking"


def test_defect_1_no_assumption_wording_in_the_document(db_conn) -> None:
    """Формулировки-предположения в документе нет."""
    data = load_report_data(FULL_INN, db_conn)
    text = "\n".join(provenance(data, "модель", datetime(2026, 1, 1)))
    assert "предположение" not in text
    assert "по умолчанию" not in text
    assert "определена формой отчётности" in text


# --- 2. Веса групп ----------------------------------------------------------


def test_defect_2_percent_formatting() -> None:
    """Доля 0,3 печатается как 30,0 %, а не как 3000,0 %."""
    assert _percent(Decimal("0.3")) == "30,0 %"
    assert _percent(Decimal("0.2")) == "20,0 %"
    assert _percent(Decimal("1")) == "100,0 %"
    assert _percent(None) == "—"


def test_defect_2_weights_are_stored_as_shares(db_conn) -> None:
    """Номинальный и фактический вес — в одних единицах, оба доли."""
    data = load_report_data(FULL_INN, db_conn)
    assert data.groups
    for group in data.groups:
        assert Decimal(0) <= group["nominal_weight"] <= Decimal(1), group["group_code"]
        assert Decimal(0) <= group["effective_weight"] <= Decimal(1), group["group_code"]
    total = sum(item["nominal_weight"] for item in data.groups)
    assert total == Decimal(1)


def test_defect_2_rendered_weights_are_sane(db_conn) -> None:
    """В таблице приложения веса не превышают ста процентов."""
    table = groups_table(load_report_data(FULL_INN, db_conn))
    assert table is not None
    for row in table.rows:
        for cell in row:
            if cell.endswith(" %"):
                value = Decimal(cell.removesuffix(" %").replace(",", "."))
                assert value <= 100, row


# --- 3. «Рассчитано» против «участвует в балле» -----------------------------


def test_defect_3_column_says_what_it_counts(db_conn) -> None:
    """Графа названа так, как считает: показатели в балле."""
    data = load_report_data(FULL_INN, db_conn)
    groups = groups_table(data)
    assert groups is not None
    assert "Показателей в балле" in groups.header
    assert "Показателей" not in groups.header

    metrics = metrics_table(data)
    assert "Участвует в балле" in metrics.header


def test_defect_3_zero_in_the_column_means_zero_in_the_score(db_conn) -> None:
    """Ноль в графе не означает, что показатель не рассчитан."""
    data = load_report_data(FULL_INN, db_conn)
    turnover = [item for item in data.groups if item["group_code"] == "turnover"]
    assert turnover, "группа «Оборачиваемость» есть в разложении"
    assert turnover[0]["metrics_used"] == 0
    calculated = [
        item
        for item in data.metrics
        if item.group_name == "Оборачиваемость" and any(item.values.values())
    ]
    assert calculated, "при нуле в графе показатели группы всё же рассчитаны"


# --- 4. Класс и отказ в классе ----------------------------------------------


def test_defect_4_stop_factor_outranks_sufficiency(db_conn) -> None:
    """При стоп-факторе класс присваивается, а не отменяется узостью основания."""
    data = load_report_data(STOPPED_INN, db_conn)
    assert data.stop_factor_code
    assert data.class_code, "стоп-фактор обязан присвоить класс"
    assert data.assessment["no_class_reason"] is None
    assert data.breadth_reason, "узость основания фиксируется отдельно"


def test_defect_4_no_contradiction_in_the_document(db_conn) -> None:
    """Документ не содержит одновременно присвоения и неприсвоения класса."""
    for inn in (FULL_INN, STOPPED_INN, NO_CLASS_INN):
        data = load_report_data(inn, db_conn)
        problems = [item.code for item in check_document(data, "")]
        assert "verdict_is_ambiguous" not in problems, inn


def test_defect_4_score_is_withheld_without_scoring(db_conn) -> None:
    """Класс от стоп-фактора при узком основании балла не раскрывает."""
    data = load_report_data(STOPPED_INN, db_conn)
    assert not data.score_in_summary
    assert not data.score_in_appendix
    assert groups_table(data) is None


# --- 5. Причины исключения --------------------------------------------------


def test_defect_5_every_exclusion_declares_its_kind() -> None:
    """У каждой причины исключения есть машинный вид."""
    checked = 0
    for metric in METRICS.metrics:
        if metric.in_scoring:
            assert metric.scoring_exclusion_kind is None, metric.code
            continue
        checked += 1
        assert metric.scoring_exclusion_kind is not None, metric.code
    assert checked >= 10


def test_defect_5_hierarchy_is_fixed() -> None:
    """Иерархия причин задана явно и в согласованном порядке."""
    assert EXCLUSION_ORDER == (
        ExclusionKind.STOP_FACTOR,
        ExclusionKind.NO_LEVEL_SCALE,
        ExclusionKind.DUPLICATE,
        ExclusionKind.NO_DATA,
    )
    assert ExclusionKind.STOP_FACTOR.rank < ExclusionKind.NO_DATA.rank


def test_defect_5_reasons_are_ordered_by_hierarchy(db_conn) -> None:
    """В приложении причины идут по иерархии, а не как попало."""
    data = load_report_data(NO_CLASS_INN, db_conn)
    ranks = [item.exclusion_rank for item in data.excluded_by_methodology]
    assert ranks == sorted(ranks)


def test_defect_5_reason_does_not_assert_a_fact_about_the_organisation() -> None:
    """Текст причины не утверждает того, чего у организации может не быть.

    У организации с капиталом 82 251 тыс. руб. документ разъяснял, чем плох
    отрицательный собственный капитал.
    """
    equity = METRICS.require("equity")
    assert equity.scoring_exclusion_kind is ExclusionKind.STOP_FACTOR
    assert "отрицательный собственный капитал опускает" not in equity.scoring_exclusion_reason


def test_defect_5_table_names_the_kind(db_conn) -> None:
    """Таблица причин называет вид причины отдельной графой."""
    table = not_calculated_table(load_report_data(NO_CLASS_INN, db_conn))
    assert table is not None
    assert "Вид причины" in table.header


# --- 6. Состав отчётности ---------------------------------------------------


def test_defect_6_quarantined_sets_are_not_called_accepted(db_conn) -> None:
    """Отбракованный комплект не перечисляется среди принятых."""
    data = load_report_data(NO_CLASS_INN, db_conn)
    assert data.quarantined_sources, "у пробы есть отбракованный комплект"
    accepted = {item["report_year"] for item in data.accepted_sources}
    quarantined = {item["report_year"] for item in data.quarantined_sources}
    assert not (accepted & quarantined)

    text = "\n".join(provenance(data, "модель", datetime(2026, 1, 1)))
    year = str(next(iter(quarantined)))
    accepted_line = next(line for line in text.split("\n") if "в расчёте:" in line)
    assert year not in accepted_line


def test_defect_6_consistency_control_catches_the_contradiction(db_conn) -> None:
    """Контроль ловит расхождение между «Ограничениями» и «Происхождением»."""
    data = load_report_data(NO_CLASS_INN, db_conn)
    year = data.quarantined_sources[0]["report_year"]

    silent = check_document(data, "- Оговорок нет")
    assert any(item.code == "quarantine_not_disclosed" for item in silent)

    honest = check_document(
        data, f"- Отчётность за период 31.12.{year} не включена в расчёт"
    )
    assert not any(item.code.startswith("source") for item in honest)


def test_defect_6_accepted_year_called_excluded_is_caught(db_conn) -> None:
    """Принятый комплект, названный невключённым, — тоже противоречие."""
    data = load_report_data(FULL_INN, db_conn)
    year = data.accepted_sources[0]["report_year"]
    problems = check_document(data, f"- Отчётность за {year} год не включена в расчёт")
    assert any(item.code == "source_both_excluded_and_accepted" for item in problems)


# --- 7. Нумерация таблиц ----------------------------------------------------


@pytest.mark.parametrize("inn", [FULL_INN, STOPPED_INN, NO_CLASS_INN])
def test_defect_7_table_titles_carry_no_numbers(inn: str, db_conn) -> None:
    """Номер в заголовке таблицы не зашит: его ставит сборщик документа."""
    data = load_report_data(inn, db_conn)
    tables = [
        metrics_table(data),
        not_calculated_table(data),
        groups_table(data),
        checks_table(data),
    ]
    for table in tables:
        if table is not None:
            assert not table.title.startswith("Таблица"), table.title


def test_defect_7_numbering_is_continuous(db_conn, tmp_path) -> None:
    """Нумерация сквозная и без пропусков даже там, где таблица не строится."""
    from docx import Document

    from finlib.llm.service import Conclusion
    from finlib.report.document import build_report

    answer = "\n".join(
        f"### {number}. {title}\nТекст."
        for number, title in (
            (2, "Фактическая база"),
            (3, "Аналитическая интерпретация"),
            (4, "Риски"),
            (5, "Ограничения анализа"),
            (6, "Вопросы к организации"),
        )
    )
    conclusion = Conclusion(
        inn=STOPPED_INN,
        report_date=date(2024, 12, 31),
        text=answer,
        model="тестовая-модель",
        attempt=1,
        checked_numbers=0,
    )
    report = build_report(
        STOPPED_INN, db_conn, conclusion=conclusion, directory=tmp_path
    )
    document = Document(report.path)
    numbers = [
        int(item.text.split()[1].rstrip("."))
        for item in document.paragraphs
        if item.text.startswith("Таблица ")
    ]
    assert numbers == list(range(1, len(numbers) + 1)), numbers
    assert len(numbers) == len(document.tables)
