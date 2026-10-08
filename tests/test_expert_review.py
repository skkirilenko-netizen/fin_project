"""Пакет экспертной проверки из истории маршрута: синтетика в findb_test.

Пакет воспроизводим: маршрут дня читается из `routing_day`, а не
пересчитывается на сегодня; та же дата и та же база дают тот же пакет.
"""

import json
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from finlib.config import settings
from finlib.db import PgConnection, execute
from finlib.normalize.lines import load_lines

sys.path.insert(0, str(settings.base_dir / "eval"))
import expert_review_run as review  # noqa: E402

DAY = date(2090, 3, 2)
INNS = [f"00000000{number:02d}" for number in range(90, 96)]
THOUSANDS = load_lines().units.name_of("384")


def _point(conn: PgConnection, inn: str, basket: str, grounds: list[str],
           debt: str | None, kind: str = "run") -> None:
    """Синтетическая точка маршрута дня в откатываемой транзакции."""
    inputs = {"metrics": {"debt_total": debt, "equity_ratio": "0.25"} if debt else {},
              "unit": THOUSANDS, "cash": "150",
              "refinance": {"due": "700", "offered": "0", "cash": "150",
                            "unit": THOUSANDS, "days": 365}}
    execute(
        "INSERT INTO routing_history (inn, as_of, kind, standard, basket, grounds, "
        "grounds_all, inputs, fingerprint, report_date) VALUES (%(inn)s, %(day)s, "
        "%(kind)s, 'rsbu', %(basket)s, %(grounds)s, %(all)s, %(inputs)s, 'fp', '2089-12-31')",
        {"inn": inn, "day": DAY, "kind": kind, "basket": basket, "grounds": grounds,
         "all": [*grounds, "refinancing_gap"], "inputs": json.dumps(inputs)},
        conn=conn,
    )


@pytest.fixture
def day(db_conn: PgConnection, monkeypatch: pytest.MonkeyPatch) -> PgConnection:
    """Маршрут дня: один в «Разборе», пять «Без внимания»; у одного крупный долг."""
    execute("DELETE FROM routing_history WHERE inn = ANY(%(inns)s)", {"inns": INNS},
            conn=db_conn)
    _point(db_conn, INNS[0], "review", ["market_spread_extreme"], "5000")
    for number, debt in zip(range(1, 6), ["100", "900", None, "500", "50"], strict=True):
        _point(db_conn, INNS[number], "clear", [], debt)
    monkeypatch.setattr(review, "cards", lambda: {
        inn: {"name_rus": f"Эмитент {inn[-2:]}"} for inn in INNS})
    monkeypatch.setattr(review, "universe", lambda: INNS)
    volumes = {INNS[5]: Decimal(10) ** 12}
    monkeypatch.setattr(review, "_outstanding", lambda inn: volumes.get(inn))
    monkeypatch.setattr(review, "SAMPLE_SIZE", 3)
    return db_conn


def _rows(path: Path, title: str) -> list[dict]:
    """Строки листа словарями по шапке."""
    sheet = load_workbook(path)[title]
    head, *rest = [list(row) for row in sheet.iter_rows(values_only=True)]
    return [dict(zip(head, row, strict=True)) for row in rest]


def test_the_package_is_built_from_the_recorded_route(day: PgConnection, tmp_path: Path) -> None:
    """Листы, корзина, решающее основание, величины с единицей и пустые графы вердикта."""
    (tmp_path / "cards").mkdir()
    (tmp_path / "cards" / f"{INNS[0]}.md").write_text("карточка", encoding="utf-8")
    out = tmp_path / "pack.xlsx"
    review.build(day, DAY, out)
    assert load_workbook(out).sheetnames == ["Сводка", "Разбор", "Без внимания", "Выборка"]
    [first] = _rows(out, "Разбор")
    assert first["ИНН"] == INNS[0] and first["корзина"] == "Разбор"
    assert first["решающее основание"]
    assert first["решающее основание"] != "market_spread_extreme", "код вместо наименования"
    assert THOUSANDS in first["debt_total"] and first["debt_total"].startswith("5\u00a0000")
    assert THOUSANDS in first["платежи 12 месяцев"]
    assert first["карточка"] == f"cards/{INNS[0]}.md"
    for column in review.VERDICT:
        assert column in first and first[column] is None
    assert len(_rows(out, "Без внимания")) == 5


def test_the_sample_takes_large_bonds_then_reported_debt(
    day: PgConnection, tmp_path: Path
) -> None:
    """Выборка: крупный долг по облигациям, затем по долгу отчётности; без долга — не берётся."""
    out = tmp_path / "pack.xlsx"
    summary = dict(review.build(day, DAY, out))
    sample = _rows(out, "Выборка")
    assert [row["ИНН"] for row in sample] == [INNS[5], INNS[2], INNS[4]]
    assert sample[0]["основание отбора"] == "крупный долг по облигациям"
    assert sample[1]["основание отбора"] == "долг по отчётности, место 1"
    assert "долг не раскрыт либо единица неизвестна у 1" in summary["Добор по долгу по отчётности"]


def test_the_same_day_gives_the_same_package(day: PgConnection, tmp_path: Path) -> None:
    """Воспроизводимость: два пакета на одну дату совпадают по содержимому листов."""
    first, second = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    review.build(day, DAY, first)
    review.build(day, DAY, second)
    for title in ("Сводка", "Разбор", "Без внимания", "Выборка"):
        assert _rows(first, title) == _rows(second, title)


def test_a_repeat_of_the_day_is_read_as_the_last_point(
    day: PgConnection, tmp_path: Path
) -> None:
    """Повтор дня после прогона по расписанию — маршрут дня по нему, как у отчёта изменений."""
    execute(
        "INSERT INTO routing_run (kind, as_of, status) VALUES ('run', %(day)s, 'done')",
        {"day": DAY}, conn=day,
    )
    execute(
        "INSERT INTO routing_history (run_id, inn, as_of, kind, standard, basket, grounds, "
        "grounds_all, inputs, fingerprint) SELECT max(id), %(inn)s, %(day)s, 'repeat', "
        "'rsbu', 'review', ARRAY['market_spread_extreme'], ARRAY['market_spread_extreme'], "
        "'{}'::jsonb, 'fp2' FROM routing_run",
        {"inn": INNS[1], "day": DAY}, conn=day,
    )
    out = tmp_path / "pack.xlsx"
    review.build(day, DAY, out)
    assert {row["ИНН"] for row in _rows(out, "Разбор")} == {INNS[0], INNS[1]}


def test_a_day_without_route_is_refused(db_conn: PgConnection, tmp_path: Path) -> None:
    """Дня в истории нет — отказ, а не пустой пакет."""
    with pytest.raises(review.NoRouteError):
        review.build(db_conn, date(2091, 1, 1), tmp_path / "pack.xlsx")
    assert not (tmp_path / "pack.xlsx").exists()


def test_an_existing_package_is_not_overwritten(tmp_path: Path, capsys) -> None:
    """Готовый файл не переписывается: отказ до обращения к базе."""
    out = tmp_path / "pack.xlsx"
    out.write_bytes(b"x")
    assert review.main(["--as-of", "2090-03-02", "--out", str(out)]) == 1
    assert "не переписывается" in capsys.readouterr().out
    assert out.read_bytes() == b"x"
