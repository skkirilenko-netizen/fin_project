"""Тесты надзорных сигналов (задача 15).

Экспертная оценка показала, что по ПК «Стройсервис» система не отразила
единственный существенный сигнал: при чистой прибыли 40 839 тыс. руб.
собственный капитал сократился с 691 до −442 тыс. руб. Выявление такого
не может оставаться на усмотрение модели.
"""

import copy
from decimal import Decimal

import pytest
import yaml
from pydantic import ValidationError

from finlib.metrics.formula import FormulaError, PrevRef, evaluate, parse_formula
from finlib.report.data import load_report_data
from finlib.scoring.signals import (
    CalibrationStatus,
    SignalLevel,
    SignalsCatalog,
    default_path,
    evaluate_signals,
    load_signals,
    revision_intensity,
    shown,
    structure_shifts,
)

CATALOG = load_signals()

# Неразрывный пробел между разрядами: в документе число не должно рваться
# по переносу строки.
NBSP = " "

# Все пороги методики: четыре сигнала-выражения и два, считаемые не формулой.
# Требования к порогу одни и те же независимо от того, как получена величина.
RULES = (*CATALOG.signals, CATALOG.structure_shift, CATALOG.revision_intensity)


def raw() -> dict:
    """Справочник сигналов в исходном виде, до разбора моделью."""
    return yaml.safe_load(default_path().read_text(encoding="utf-8"))

# Отчётность ПК «Стройсервис»: тот самый случай из экспертной оценки.
NOW = {"1300": Decimal(-442), "2400": Decimal(40839), "2110": Decimal(44771), "1600": Decimal(418)}
BEFORE = {"1300": Decimal(691), "2400": Decimal(4628), "2110": Decimal(4920), "1600": Decimal(916)}


def codes(current: dict, previous: dict) -> set[str]:
    """Коды сработавших сигналов."""
    return {item.code for item in evaluate_signals(current, previous, CATALOG)}


# --- методика ----------------------------------------------------------------


def test_every_threshold_declares_its_origin() -> None:
    """Порог без происхождения неотличим от выдуманного."""
    for signal in CATALOG.signals:
        assert signal.origin.strip(), signal.code
    assert CATALOG.structure_shift.origin.strip()
    assert CATALOG.revision_intensity.origin.strip()


def test_every_signal_has_a_prescribed_wording() -> None:
    """Формулировка задана методикой, а не сочиняется моделью."""
    for signal in CATALOG.signals:
        assert "{value}" in signal.text, signal.code
        assert len(signal.text) > 80, signal.code


def test_every_threshold_declares_its_calibration_status() -> None:
    """Зрелость порога объявлена у всех шести, а не у одного.

    Происхождение и зрелость — разные сведения: origin отвечает, откуда
    взялась величина, статус — на скольких наблюдениях она проверена.
    Пока наблюдений три, и предварительны все.
    """
    assert len(RULES) == 6
    for rule in RULES:
        assert rule.calibration_status.strip(), rule.name
        assert rule.calibration is CalibrationStatus.PRELIMINARY, rule.name
        assert rule.preliminary, rule.name
        assert "3 организациям" in rule.calibration_status, rule.name


def test_calibration_status_is_machine_readable() -> None:
    """Статус начинается признаком из перечня, а не свободным текстом.

    Без машинного признака «предварительный порог» отличался бы от
    установленного только на глаз, и переход к калиброванным (задача 19)
    нечем было бы проверить.
    """
    payload = raw()
    payload["signals"][0]["calibration_status"] = "проверено на практике"
    with pytest.raises(ValidationError, match="статус калибровки"):
        SignalsCatalog(**payload)


def test_every_rule_declares_its_condition() -> None:
    """Формулировка без условия срабатывания методику не проходит."""
    for rule in RULES:
        assert rule.condition is not None, rule.name


@pytest.mark.parametrize("place", ["structure_shift", "revision_intensity"])
def test_condition_is_required_for_rules_computed_in_code(place: str) -> None:
    """Условие обязательно и там, где величина считается не формулой.

    Иначе отсечка лежала бы в методике, а знак сравнения — в коде, и по
    справочнику нельзя было бы сказать, когда печатается формулировка.
    """
    payload = raw()
    del payload[place]["condition"]
    with pytest.raises(ValidationError, match="condition"):
        SignalsCatalog(**payload)


def test_condition_is_required_for_signal_expressions() -> None:
    """То же для сигналов, задаваемых выражением."""
    payload = raw()
    del payload["signals"][0]["condition"]
    with pytest.raises(ValidationError, match="condition"):
        SignalsCatalog(**payload)


def test_structure_shift_follows_the_declared_condition() -> None:
    """Сравнение идёт по объявленному условию, а не по зашитому в коде."""
    payload = copy.deepcopy(raw())
    payload["structure_shift"]["condition"] = "gt"
    strict = SignalsCatalog(**payload)
    exactly = strict.structure_shift.threshold_points
    shares_now = {"1200": Decimal(50) + exactly}
    shares_before = {"1200": Decimal(50)}
    names = {"1200": "Оборотные активы"}
    # Ровно на пороге: gte даёт сигнал, gt молчит.
    assert structure_shifts(shares_now, shares_before, names, CATALOG)
    assert structure_shifts(shares_now, shares_before, names, strict) == []


def test_levels_are_declared() -> None:
    """У каждого сигнала объявлен вес."""
    levels = {signal.level for signal in CATALOG.signals}
    assert levels <= {SignalLevel.ATTENTION, SignalLevel.SUPERVISORY}
    assert SignalLevel.SUPERVISORY in levels


# --- язык выражений ----------------------------------------------------------


def test_prev_is_part_of_the_formula_language() -> None:
    """prev(код) разбирается тем же интерпретатором, без eval."""
    tree = parse_formula("1300 - prev(1300)")
    assert evaluate(tree, {"1300": Decimal(-442)}, {"1300": Decimal(691)}, {}) == Decimal(-1133)


def test_prev_requires_the_previous_period() -> None:
    """Без предыдущего периода выражение не вычисляется, а не подменяется нулём."""
    with pytest.raises(FormulaError, match="нет предыдущего периода"):
        evaluate(parse_formula("prev(1300)"), {"1300": Decimal(1)}, None, {})


def test_prev_requires_the_previous_period_to_be_disclosed() -> None:
    """Нераскрытая строка за предыдущий период — не ноль."""
    with pytest.raises(FormulaError, match="не раскрыта"):
        evaluate(parse_formula("prev(1300)"), {"1300": Decimal(1)}, {"1300": None}, {})


def test_numeric_literals_stay_forbidden() -> None:
    """Запрет числовых литералов новая функция не ослабила."""
    with pytest.raises(FormulaError, match="магическое"):
        parse_formula("1300 - 500")


def test_prev_counts_as_needing_the_previous_period() -> None:
    """Показатель с prev() требует предыдущего периода наравне с avg()."""
    from finlib.metrics.formula import average_codes

    assert average_codes(parse_formula("1300 - prev(1300)")) == {"1300"}
    assert isinstance(parse_formula("prev(1600)"), PrevRef)


# --- срабатывание на реальном случае -----------------------------------------


def test_equity_withdrawal_is_detected() -> None:
    """Изъятие капитала в пользу участников — тот сигнал, которого не хватало."""
    assert "equity_withdrawal" in codes(NOW, BEFORE)


def test_withdrawal_message_names_both_values() -> None:
    """Формулировка называет и расхождение, и финансовый результат."""
    hit = next(
        item for item in evaluate_signals(NOW, BEFORE, CATALOG) if item.code == "equity_withdrawal"
    )
    assert "41" in hit.message and "972" in hit.message
    assert "40" in hit.message and "839" in hit.message
    assert hit.level is SignalLevel.SUPERVISORY
    # Знак выражен словами, минус в величине читался бы как опечатка.
    assert "-41" not in hit.message


def test_transit_structure_is_detected() -> None:
    """Оборот в сто раз выше валюты баланса — признак транзитной структуры."""
    assert "transit_structure" in codes(NOW, BEFORE)


def test_extreme_margin_is_detected() -> None:
    """Чистая рентабельность 91 % требует объяснения."""
    assert "extreme_margin" in codes(NOW, BEFORE)


def test_ordinary_organisation_triggers_nothing() -> None:
    """На обычных величинах сигналы молчат."""
    current = {
        "1300": Decimal(1200),
        "2400": Decimal(200),
        "2110": Decimal(5000),
        "1600": Decimal(3000),
    }
    previous = {
        "1300": Decimal(1000),
        "2400": Decimal(180),
        "2110": Decimal(4800),
        "1600": Decimal(2900),
    }
    assert codes(current, previous) == set()


def test_contribution_is_the_mirror_of_withdrawal() -> None:
    """Прирост капитала выше прибыли — признак взносов, а не распределения."""
    current = {
        "1300": Decimal(5000),
        "2400": Decimal(100),
        "2110": Decimal(4000),
        "1600": Decimal(6000),
    }
    previous = {
        "1300": Decimal(1000),
        "2400": Decimal(90),
        "2110": Decimal(3800),
        "1600": Decimal(5000),
    }
    found = codes(current, previous)
    assert "equity_contribution" in found
    assert "equity_withdrawal" not in found


def test_missing_data_does_not_trigger() -> None:
    """Нераскрытая строка сигнала не даёт: подстановки нуля нет."""
    assert codes({"1300": Decimal(-442), "2400": None}, BEFORE) == set()


# --- структурный сдвиг и пересмотр -------------------------------------------


def test_structure_shift_is_detected() -> None:
    """Перераспределение пятой части баланса — событие, а не колебание."""
    found = structure_shifts(
        {"1300": Decimal(-105), "1200": Decimal(50)},
        {"1300": Decimal(75), "1200": Decimal(45)},
        {"1300": "Капитал и резервы", "1200": "Оборотные активы"},
        CATALOG,
    )
    assert [item.details["line_code"] for item in found] == ["1300"]
    assert "Капитал и резервы" in found[0].message


def test_structure_shift_ignores_small_moves() -> None:
    """Колебание доли сигналом не считается."""
    assert (
        structure_shifts(
            {"1200": Decimal(50)}, {"1200": Decimal(45)}, {"1200": "Оборотные"}, CATALOG
        )
        == []
    )


@pytest.mark.parametrize(("mismatches", "sets"), [(28, 5), (52, 4)])
def test_revision_intensity_covers_both_observations(mismatches: int, sets: int) -> None:
    """Обе наблюдаемые организации отсечку проходят: 5,6 и 13,0 на комплект."""
    hit = revision_intensity(mismatches, sets, CATALOG)
    assert hit is not None
    assert str(mismatches) in hit.message


def test_revision_intensity_stays_silent_below_the_threshold() -> None:
    """Редкий пересмотр сигналом не считается."""
    assert revision_intensity(4, 4, CATALOG) is None
    assert revision_intensity(10, 0, CATALOG) is None


# --- расчёт и документ -------------------------------------------------------


def test_signals_are_computed_and_stored(db_conn) -> None:
    """Сигналы считаются в оценке и доходят до документа."""
    data = load_report_data("2100010824", db_conn)
    found = {item["signal_code"] for item in data.signals}
    assert "equity_withdrawal" in found
    assert "transit_structure" in found
    assert all(item["message"].strip() for item in data.signals)


def test_supervisory_signals_come_first(db_conn) -> None:
    """Надзорные сигналы стоят впереди требующих внимания."""
    data = load_report_data("2100010824", db_conn)
    levels = [item["level"] for item in data.signals]
    assert levels == sorted(levels, key=lambda item: item != "supervisory")


def test_signals_reach_the_document_without_the_model(db_conn, tmp_path) -> None:
    """Сигналы детерминированы: в справке без текстовой части они остаются."""
    from docx import Document

    from finlib.report.document import SIGNALS_TITLE, build_report

    report = build_report("2100010824", db_conn, directory=tmp_path, with_text=False)
    text = "\n".join(item.text for item in Document(report.path).paragraphs)
    assert SIGNALS_TITLE in text
    assert "Признаки распределения в пользу участников" in text
    assert "не объясняется финансовым результатом" in text


def test_signal_wording_is_not_paraphrased(db_conn, tmp_path) -> None:
    """В документ идёт формулировка методики, а не пересказ."""
    from docx import Document

    from finlib.report.document import build_report

    data = load_report_data("2100010824", db_conn)
    report = build_report("2100010824", db_conn, directory=tmp_path, with_text=False)
    text = "\n".join(item.text for item in Document(report.path).paragraphs)
    for signal in data.signals:
        assert signal["message"] in text, signal["signal_code"]


def test_signal_value_is_written_in_russian_form() -> None:
    """Разряды разделены неразрывным пробелом, десятичный знак — запятая.

    Разрядность берётся у сигнала, а не у места вывода: кратность с одним
    знаком, доля с двумя, денежная величина целыми тысячами.
    """
    assert shown(CATALOG.rule_for("equity_withdrawal"), Decimal("-41972.000")) == (
        f"41{NBSP}972"
    )
    assert shown(CATALOG.rule_for("transit_structure"), Decimal("107.1084")) == "107,1"
    assert shown(CATALOG.rule_for("extreme_margin"), Decimal("0.9122")) == "0,91"
    assert shown(CATALOG.rule_for("structure_shift_1300"), Decimal("-181.1783")) == (
        "181,2"
    )


def test_negative_value_keeps_the_typographic_minus() -> None:
    """Знак у величины — математический минус, а не дефис.

    Величина стоит в документе среди прозы, где дефис читается как тире,
    а тире перед числом минусом не считается.
    """
    rule = CATALOG.rule_for("equity_contribution")
    assert not rule.as_absolute
    assert shown(rule, Decimal(-5)) == "−5"


def test_unknown_signal_code_has_no_rule() -> None:
    """Кода нет в справочнике — разрядности взять неоткуда, и это видно."""
    assert CATALOG.rule_for("no_such_signal") is None


def test_absolute_value_is_declared_by_the_methodology() -> None:
    """Подстановка по модулю — решение методики, а не приём кода.

    У структурного сдвига направление читается из долей на начало и на конец
    периода, и модуль объявлен в справочнике, а не зашит в расчёт.
    """
    assert raw()["structure_shift"]["as_absolute"] is True
    assert CATALOG.structure_shift.as_absolute


def test_preliminary_threshold_is_marked_as_such() -> None:
    """Порог, подогнанный под известный ответ, помечен предварительным.

    Отсечка интенсивности пересмотра подобрана так, чтобы сработали обе
    наблюдаемые организации. На выборке из трёх это подгонка, а не
    калибровка, и признак не даёт выдать одно за другое.
    """
    rule = CATALOG.revision_intensity
    assert rule.preliminary
    assert "не калибровка" in rule.origin
