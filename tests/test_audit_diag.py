"""Диагностика расхождений маршрута: правила разбора на синтетике, без базы и сети."""

import json
import sys
from collections.abc import Iterator
from contextlib import contextmanager
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from finlib.config import settings
from finlib.scoring.routing import Finding, load_routing
from finlib.sources import cbonds, moex
from finlib.sources.cbonds_events import Guarantee, Issue
from finlib.sources.market import Market, Point

sys.path.insert(0, str(settings.base_dir / "eval"))

import audit_diag as diag  # noqa: E402


def _state(**over: object) -> diag.IsinState:
    """Бумага в обращении, сопоставленная, без строк в окне."""
    found = diag.IsinState(isin="RU000A", status="в обращении", maturity=None, mapped=True)
    for key, value in over.items():
        setattr(found, key, value)
    return found


def test_focus_is_matched_by_bare_name_and_ambiguity_is_said() -> None:
    """«Газпром» опознаётся ровно одной карточкой; неоднозначность называет кандидатов."""
    known = {
        "1": {"name_rus": "ПАО «Газпром»"},
        "2": {"name_rus": "Газпром Капитал ООО"},
        "3": {"name_rus": "Газпром нефть"},
        "4": {"name_rus": "ЛОЭСК"},
        "5": {"name_rus": "ЛОЭСК-Сети"},
    }
    found = diag.focus_inns(known, ("Газпром", "ПИК", "ЛОЭСК"), {})
    assert found["Газпром"] == (["1"], [])
    assert found["ПИК"] == ([], [])
    assert found["ЛОЭСК"] == (["4"], [])
    known["6"] = {"name_rus": "АО «ЛОЭСК»"}
    assert diag.focus_inns(known, ("ЛОЭСК",), {})["ЛОЭСК"] == ([], ["4", "6"])
    # Явное указание снимает неоднозначность.
    assert diag.focus_inns(known, ("ЛОЭСК",), {"ЛОЭСК": "6"})["ЛОЭСК"] == (["6"], [])
    # Не опознанный фокус попадает на лист строкой, а не исчезает.
    rows = diag.focus_rows("card-status", {"ЛОЭСК": ([], ["4", "6"])})
    assert rows and rows[0][1] == "ЛОЭСК" and rows[0][-1] == "кандидаты: 4, 6"
    assert len(rows[0]) == len(diag.HEADERS["card-status"])


def test_gap_reason_goes_from_common_to_particular() -> None:
    """Причина молчания — первая подходящая: доставка, погашение, ISIN, срезы, сделки, цена."""
    assert diag.gap_reason([_state()], 12, 10, 0).startswith("сбой загрузки")
    repaid = [_state(status="погашена")]
    assert diag.gap_reason(repaid, 0, 10, 0) == "выпусков в обращении нет (погашена)"
    assert diag.gap_reason([_state(isin="")], 0, 10, 0) == "у выпусков в обращении нет ISIN"
    assert diag.gap_reason([_state(mapped=False)], 0, 10, 0).startswith("ISIN не сопоставлен")
    assert diag.gap_reason([_state()], 0, 10, 0).startswith("бумаг нет в срезах")
    assert diag.gap_reason([_state(rows=5)], 0, 10, 0) == "в срезах есть, сделок нет"
    refused = _state(rows=5, traded=3, refused={"structural"})
    assert diag.gap_reason([refused], 0, 10, 0).endswith("structural")
    priced = _state(rows=5, traded=3, priced=3)
    assert diag.gap_reason([priced], 0, 10, 2).startswith("дней без ориентира в окне: 2")
    assert "ряд не пересобран" in diag.gap_reason([priced], 0, 10, 0)


def test_coverage_names_why_the_guarantor_is_not_taken() -> None:
    """Покрытие: не доставлено, не тот вид, без ИНН, вне списка, SPV не взяла корзину."""
    accepted = ("Поручитель", "Гарант")
    backing = [{"status_name_rus": "Поручитель", "guarantor_inn": "7736050003"}]
    assert diag.coverage_verdict(True, [], False, accepted, set(), False).startswith(
        "поручительства не доставлены"
    )
    assert diag.coverage_verdict(True, [], True, accepted, set(), False) == (
        "поручителя нет у источника"
    )
    assert diag.coverage_verdict(False, [], True, accepted, set(), False) == ""
    offer = [{"status_name_rus": "Оферент", "guarantor_inn": "1"}]
    assert "Оферент" in diag.coverage_verdict(True, offer, True, accepted, {"1"}, False)
    nameless = [{"status_name_rus": "Гарант", "guarantor_inn": ""}]
    assert "без ИНН" in diag.coverage_verdict(True, nameless, True, accepted, set(), False)
    assert diag.coverage_verdict(True, backing, True, accepted, set(), False) == (
        "поручитель вне списка маршрута"
    )
    listed = {"7736050003"}
    assert diag.coverage_verdict(True, backing, True, accepted, listed, False).startswith("SPV")
    assert diag.coverage_verdict(True, backing, True, accepted, listed, True) == ""
    # У обычного эмитента корзина поручителя не переносится — и это не расхождение.
    assert diag.coverage_verdict(False, backing, True, accepted, listed, False) == ""


def test_volumes_compare_route_issues_and_the_source_list() -> None:
    """Непогашенные выпуски вне статусов маршрута и амортизация дают расхождение; валюта названа."""
    issues = [
        {"isin_code": "RU1", "status_name_rus": "В обращении", "outstanding_volume": "1000",
         "nominal_price": "1000", "outstanding_nominal_price": "500", "currency_name": "RUB"},
        {"isin_code": "RU2", "status_name_rus": "Дефолт", "outstanding_volume": "200"},
        {"isin_code": "RU3", "status_name_rus": "Погашена", "outstanding_volume": "900"},
        {"isin_code": "RU4", "status_name_rus": "В обращении", "outstanding_volume": None,
         "currency_name": "CNY"},
        {"isin_code": "RU5", "status_name_rus": "В обращении", "outstanding_volume": "100",
         "currency_name": "USD"},
        # Досрочно погашенный, аннулированный и планируемый — не непогашенные
        # (разбор digest 09.10.2026): ни в сумму, ни в «без объёма» не идут.
        {"isin_code": "RU6", "status_name_rus": "Досрочно погашена", "outstanding_volume": "5000"},
        {"isin_code": "", "status_name_rus": "Аннулирована", "outstanding_volume": None},
        {"isin_code": "", "status_name_rus": "Планируется", "outstanding_volume": None},
    ]
    universe = [{"isin_code": "RU1", "outstanding_volume": "1000"},
                {"isin_code": "RU9", "outstanding_volume": "50"}]
    got = diag.volumes_of(issues, universe, Decimal(1100))
    assert got.by_issues == Decimal(1300) and got.amortized == Decimal(800)
    assert got.universe == Decimal(1050) and got.without_volume == 1
    assert got.foreign == ["RU5: USD"]
    assert got.only_issuer == ["RU4", "RU5"] and got.only_universe == ["RU9"]
    assert diag.share_gap(got.route, got.by_issues) == Decimal(200) / Decimal(1100)
    assert diag.share_gap(None, Decimal(1)) is None
    assert diag.share_gap(Decimal(0), Decimal(0)) is None
    assert diag.share_gap(Decimal(0), Decimal(5)) == Decimal(1)
    empty = diag.volumes_of([], [], None)
    assert empty.by_issues is None and empty.universe is None


def test_card_status_against_the_basket() -> None:
    """Недействующий статус при «Без внимания», банкротство вне «Разбора», выход с преемником."""
    routing = load_routing()
    universe = routing.universe
    code = {name: key for key, name in universe.statuses.items()}
    active = code["действующая"]
    queue = universe.unconfirmed_to
    assert diag.card_mismatch(active, "clear", routing, False) == ""
    # Очередь статуса у действующего — расхождение, только если её назвал статус:
    # давность отчётности у действующего штатна (разбор digest 09.10.2026).
    assert diag.card_mismatch(active, queue, routing, False, ("reporting_two_cycles_old",)) == ""
    assert "неподтверждённый статус" in diag.card_mismatch(
        active, queue, routing, False, ("status_not_confirmed",)
    )
    assert "в маршруте эмитента нет" in diag.card_mismatch(active, None, routing, False)
    reorganized = code["в процессе реорганизации"]
    assert diag.card_mismatch(reorganized, "clear", routing, False).endswith("«Без внимания»")
    assert diag.card_mismatch(reorganized, "attention", routing, False) == ""
    bankrupt = next(key for name, key in code.items() if "банкрот" in name)
    assert diag.card_mismatch(bankrupt, "attention", routing, False).endswith("attention")
    assert diag.card_mismatch(bankrupt, "review", routing, False) == ""
    gone = universe.exclude_statuses[0]
    assert diag.card_mismatch(gone, None, routing, True) == ""
    assert "с преемником" in diag.card_mismatch(gone, "clear", routing, True)
    assert diag.card_mismatch(gone, queue, routing, False) == ""
    assert diag.card_mismatch("99", queue, routing, False) == ""
    assert "не опознан" in diag.card_mismatch("99", "clear", routing, False)
    # Без карточки — не расхождение: эмитент в маршруте по отчётности.
    assert diag.card_mismatch(None, "clear", routing, False) == ""


def test_the_book_has_a_summary_and_a_sheet_per_command(tmp_path: Path) -> None:
    """Сводка первой, лист на подкоманду с заголовком; Decimal пишется строкой."""
    out = tmp_path / "audit.xlsx"
    rows = [["1", "Эмитент", "", Decimal("1.50"), None, None, None, None, None, "да",
             "", 0, "", "", ""]]
    diag.write_book(out, {"outstanding": rows}, [("Диагностика", f"{date(2026, 10, 9):%d.%m.%Y}")])
    book = load_workbook(out)
    assert book.sheetnames == ["Сводка", "outstanding"]
    sheet = book["outstanding"]
    assert [cell.value for cell in sheet[1]] == list(diag.HEADERS["outstanding"])
    assert sheet.cell(row=2, column=4).value == "1.50"


def _issue(isin: str, status: str = "в обращении") -> Issue:
    """Выпуск эмитента с ISIN и статусом."""
    return Issue(
        emission_id=isin, name=isin, isin=isin, status=status, default=False,
        unsettled=False, maturity=None, offer=None, outstanding=None, updated=None,
    )


def test_market_gaps_end_to_end_on_disk_fixtures(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Срезы на диске, ряд и выпуски: разрыв назван причиной, фокус без разрыва — тоже строкой."""
    days = [date(2026, 10, 1) + timedelta(days=number) for number in range(12)]
    days = [day for day in days if day.weekday() < 5]
    for day in days:
        rows = [{"SECID": "RU_TRADED", "NUMTRADES": 5, "CLOSE": 99}]
        rows.append({"SECID": "RU_SILENT", "NUMTRADES": 0, "CLOSE": None})
        (tmp_path / f"xsec_{day}.json").write_text(json.dumps({"history": rows}))
    monkeypatch.setattr(moex, "CACHE", tmp_path)
    point = Point(day=days[-1], spread=Decimal(100), price=Decimal(99), weight=Decimal(1))
    market = Market(
        benchmark={day: Decimal(100) for day in days},
        issuers={"1": {days[-1]: point}},
        counted={}, census={"2": {"rows": 9, "with_price": 0, "with_spread": 0}},
        universe=2, with_isin=2,
    )
    monkeypatch.setattr(diag, "series", lambda: market)
    monkeypatch.setattr(diag, "holders", lambda: {"RU_TRADED": "1", "RU_SILENT": "2"})
    issues = {
        "1": (_issue("RU_TRADED"),),
        "2": (_issue("RU_SILENT"),),
        "3": (_issue("RU_GONE", "погашена"),),
    }
    monkeypatch.setattr(diag, "issues_of", lambda inn: (issues[inn], True))
    route = [{"inn": inn, "basket": "clear"} for inn in ("1", "2", "3")]
    rows, detail, summary = diag.market_gaps(route, {}, {"1"}, 10, days[-1])
    by_inn = {row[0]: row for row in rows}
    assert by_inn["1"][8] == "разрыва нет" and by_inn["1"][5] == 0
    assert by_inn["2"][5] == "точек нет" and by_inn["2"][8] == "в срезах есть, сделок нет"
    # Без выпусков в обращении — отдельная строка сводки, а не разрыв.
    assert by_inn["3"][8] == "без выпусков в обращении (погашена)"
    assert summary["с разрывом рынка"] == 1
    assert summary["без выпусков в обращении (не разрыв)"] == 1
    assert not any(key.startswith("причина: выпусков") for key in summary)
    assert {row[1] for row in detail} == {"RU_TRADED", "RU_SILENT", "RU_GONE"}


def test_spv_coverage_reads_the_live_rows(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """SPV с поручителем в списке и корзиной поручителя — учтено; без файла — названо."""
    (tmp_path / "guarantors_10.json").write_text(json.dumps({"items": [
        {"status_name_rus": "Поручитель", "guarantor_inn": "20",
         "guarantor_name_rus": "Поручитель"},
    ]}))
    monkeypatch.setattr(cbonds, "CACHE", tmp_path)
    led = Finding("financing_structure", "Головная", "SPV группы: корзина поручителя")
    rows = [
        SimpleNamespace(inn="10", name="SPV",
                        guarantees=(Guarantee("20", "Головная", "Поручитель", "x"),),
                        # Отметку о поручителе `led_by_guarantor` кладёт
                        # в справочные, а не в сработавшие (разбор 09.10.2026).
                        verdict=SimpleNamespace(basket="clear", findings=(), notes=(led,))),
        SimpleNamespace(inn="20", name="Головная", guarantees=(),
                        verdict=SimpleNamespace(basket="clear", findings=(), notes=())),
        SimpleNamespace(inn="30", name="SPV без файла", guarantees=(),
                        verdict=SimpleNamespace(basket="status_unknown", findings=(),
                                                notes=())),
    ]

    @contextmanager
    def fake() -> Iterator[SimpleNamespace]:
        yield SimpleNamespace(rollback=lambda: None)

    monkeypatch.setattr(diag, "connection", fake)
    monkeypatch.setattr(diag, "routing_rows", lambda conn, today: (rows, {}))
    known = {"10": {"emitent_spv": "1"}, "30": {"emitent_spv": 1}}
    found, counts = diag.spv_coverage(known, {"10"}, load_routing(), date(2026, 10, 9))
    by_inn = {row[0]: row for row in found}
    assert by_inn["10"][-1] == "учтено" and by_inn["10"][8] == "да"
    assert by_inn["30"][-1].startswith("поручительства не доставлены")
    assert "20" not in by_inn and counts["покрытие не учтено"] == 1


def test_outstanding_and_card_status_end_to_end(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Объём: выше допуска — строка; статус: реорганизация при «Без внимания» — строка."""
    (tmp_path / "emissions_1.json").write_text(json.dumps({"items": [
        {"isin_code": "RU1", "status_name_rus": "В обращении", "outstanding_volume": "1000"},
        {"isin_code": "RU2", "status_name_rus": "Дефолт", "outstanding_volume": "500"},
    ]}))
    monkeypatch.setattr(cbonds, "CACHE", tmp_path)
    monkeypatch.setattr(diag, "outstanding_universe", lambda: [
        {"emitent_inn": "1", "isin_code": "RU1", "outstanding_volume": "1000"},
    ])
    monkeypatch.setattr(diag, "_outstanding", lambda inn: Decimal(1000) if inn == "1" else None)
    route = [{"inn": "1", "basket": "clear"}, {"inn": "2", "basket": "clear"}]
    routing = load_routing()
    rows, counts = diag.outstanding(route, {}, {"2"}, Decimal("0.05"), routing)
    by_inn = {row[0]: row for row in rows}
    assert by_inn["1"][9] == "да" and by_inn["1"][7] == Decimal("0.5")
    assert by_inn["2"][9] == "сравнивать нечего" and by_inn["2"][-1].startswith("выпусков")
    code = {name: key for key, name in routing.universe.statuses.items()}
    known = {
        "1": {"name_rus": "ЛОЭСК", "emitent_statuses_id": code["в процессе реорганизации"]},
        "2": {"name_rus": "Живой", "emitent_statuses_id": code["действующая"]},
    }
    monkeypatch.setattr(diag, "bond_issuers", lambda: {"1": "ЛОЭСК", "3": "Без карточки"})
    found, counts = diag.card_status(route, known, {"2"}, routing)
    by_inn = {row[0]: row for row in found}
    assert by_inn["1"][-1].endswith("«Без внимания»")
    assert by_inn["2"][-1] == "согласовано"
    assert "3" not in by_inn and counts["расхождений"] == 1
