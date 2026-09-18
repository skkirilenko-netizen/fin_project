"""Тесты определения разделителя разрядов документа МСФО.

Отдельное число конвенции не выдаёт: «663,888» — это 663 888 при запятой
в роли разделителя разрядов и 663,888 при запятой в роли десятичного знака.
Прочтения различаются в тысячу раз, и ни один контроль сходимости ошибки
не поймает: баланс сойдётся, разделы сойдутся, коэффициенты будут верны.
"""

from decimal import Decimal

import pytest

from finlib.sources.ifrs_numbers import (
    Grouping,
    GroupingUndetermined,
    check_plausibility,
    decisive_evidence,
    detect_grouping,
    load_parsing_policy,
    parse_amount,
)

# Выдержка вёрстки ФосАгро: разряды отделены запятой, дробная часть — точкой.
PHOSAGRO = """
Итого активы 663,888 452,110
Выручка 507,718 469,004
Основные средства 1,234,567
Прибыль на акцию 12.5
Рентабельность 0.31
"""

# Русская вёрстка: разряды пробелом, дробная часть запятой.
RUSSIAN_DOC = """
Итого активы 663 888 452 110
Выручка 507 718
Прибыль на акцию 12,5
Рентабельность 0,31
"""

# Неоднозначный документ: единственная разметка — «X,YYY», и она допускает
# оба прочтения.
AMBIGUOUS_DOC = """
Итого активы 663,888
Выручка 507,718
Капитал 155,770
"""


# --- определение конвенции ----------------------------------------------------


def test_phosagro_reads_comma_as_thousands_separator() -> None:
    """У ФосАгро «663,888» — это 663 888, а не 663,888.

    Свидетельства в пользу английской конвенции: число с двумя запятыми
    («1,234,567») и дробные части после точки, где цифр не три.
    """
    found = detect_grouping(PHOSAGRO)
    assert found.convention is Grouping.ENGLISH
    assert found.english_evidence >= load_parsing_policy().digit_grouping.min_evidence
    assert found.russian_evidence == 0
    assert parse_amount("663,888", found.convention) == Decimal(663888)
    assert parse_amount("12.5", found.convention) == Decimal("12.5")


def test_handful_of_evidence_does_not_outvote_ambiguity() -> None:
    """Горстка улик против сотни неоднозначных чисел ничего не решает.

    У ФосАгро улик шесть против ста пятидесяти трёх, и все шесть ложные:
    слипшаяся строка «5 573,628 507,689» читается как число с пробелом
    между разрядами. Голосование объявляло русскую конвенцию, и величины
    расходились с отчётностью в тысячу раз.
    """
    document = "Итого активы 1 234 567\n" * 3 + "Выручка 663,888 507,718\n" * 40
    found = detect_grouping(document)
    assert not found.determined
    assert found.reason is GroupingUndetermined.OUTWEIGHED


def test_number_with_both_separators_decides() -> None:
    """Число с обоими разделителями сразу — улика бесспорная.

    «11,266.5» русской конвенцией не читается никак: один и тот же знак
    не бывает в одном числе и разрядным, и десятичным. У ФосАгро такие числа
    стоят в таблице дивидендов, которую голосование из выборки исключает.
    """
    assert decisive_evidence("дивиденд 11,266.5 млн руб.") == (0, 1)
    assert decisive_evidence("дивиденд 11.266,5 млн руб.") == (1, 0)
    assert decisive_evidence("Итого активы 663,888") == (0, 0)


def test_russian_document_is_not_read_as_english() -> None:
    """Документ с пробелом как разделителем запятой не трактуется.

    Обратный случай к ФосАгро: здесь «12,5» — двенадцать с половиной,
    а не двенадцать тысяч пятьсот.
    """
    found = detect_grouping(RUSSIAN_DOC)
    assert found.convention is Grouping.RUSSIAN
    assert found.english_evidence == 0
    assert parse_amount("663 888", found.convention) == Decimal(663888)
    assert parse_amount("12,5", found.convention) == Decimal("12.5")


def test_ambiguous_document_is_refused() -> None:
    """Документ, где осмысленны оба прочтения, не разбирается вовсе.

    Выбор по умолчанию здесь означал бы ошибку в тысячу раз у половины
    документов — молча и без единого сработавшего контроля.
    """
    found = detect_grouping(AMBIGUOUS_DOC)
    assert not found.determined
    assert found.reason is GroupingUndetermined.AMBIGUOUS
    assert found.ambiguous >= 3


def test_conflicting_conventions_are_refused() -> None:
    """Числа обеих конвенций в одном документе — отказ, а не голосование."""
    mixed = PHOSAGRO + RUSSIAN_DOC
    found = detect_grouping(mixed)
    assert not found.determined
    assert found.reason is GroupingUndetermined.CONFLICTING
    assert found.russian_evidence and found.english_evidence


def test_single_evidence_is_not_a_convention() -> None:
    """Одно свидетельство — опечатка составителя, а не конвенция документа."""
    thin = "Итого активы 1 234 567\nПрочее 42\n"
    found = detect_grouping(thin)
    assert not found.determined
    assert found.reason is GroupingUndetermined.INSUFFICIENT


def test_stray_number_of_another_convention_does_not_refuse_the_document() -> None:
    """Порог улик один на обе стороны: опечатка не отменяет конвенцию.

    Прежде победителю требовались три улики, а отказ наступал от одной
    улики против — правило требовало единогласия, которого не даёт ни одна
    вёрстка. Проверено на Акроне: 162 улики за русскую конвенцию и ноль
    против, одно подставленное число английской разметки отказывало
    документ целиком.
    """
    stray = RUSSIAN_DOC + "\nПрочие обязательства 1,234,567\n"
    found = detect_grouping(stray)
    assert found.convention is Grouping.RUSSIAN
    assert found.english_evidence == 1
    # Отброшенная улика названа: строку с ней человек обязан увидеть.
    assert found.english_samples == ("1,234,567",)


def test_three_numbers_on_each_side_are_two_conventions() -> None:
    """Три улики с каждой стороны — это уже две конвенции, и это отказ."""
    mixed = RUSSIAN_DOC + "\n".join(
        f"Прочее {i} 1,234,56{i}" for i in range(3)
    )
    found = detect_grouping(mixed)
    assert not found.determined
    assert found.reason is GroupingUndetermined.CONFLICTING


def test_evidence_names_the_numbers_behind_the_verdict() -> None:
    """Решение читается вместе с числами, на которых построено.

    Счётчик отвечает «сколько улик», а разбираться приходится с «какие»:
    «1.5» — улика за английскую конвенцию ровно до тех пор, пока не видно,
    что это ставка процента.
    """
    found = detect_grouping(RUSSIAN_DOC)
    lines = found.evidence()
    assert lines and any("за русскую" in item for item in lines)


def test_document_without_separators_needs_no_convention() -> None:
    """Где разделителей нет вовсе, выбирать не из чего."""
    plain = "Итого активы 663888\nВыручка 507718\nКапитал 155770\n"
    found = detect_grouping(plain)
    assert found.convention is Grouping.PLAIN
    assert parse_amount("663888", found.convention) == Decimal(663888)


def test_detection_counts_what_it_looked_at() -> None:
    """Определение сообщает не только решение, но и число просмотренных чисел.

    Ноль свидетельств против ноля просмотренных чисел — разные вещи,
    и по журналу их надо различать.
    """
    found = detect_grouping(PHOSAGRO)
    assert found.numbers_seen > 0
    assert "чисел просмотрено" in found.describe()
    assert found.samples


# --- разбор чисел -------------------------------------------------------------


def test_convention_is_required_to_parse() -> None:
    """Число без конвенции не читается: значения по умолчанию нет.

    Умолчание здесь означало бы ровно то, ради отказа от чего написан
    весь модуль, — угадывание.
    """
    with pytest.raises(TypeError):
        parse_amount("663,888")  # type: ignore[call-arg]


def test_same_token_reads_differently_by_convention() -> None:
    """Одна и та же запись даёт величины, различающиеся в тысячу раз."""
    assert parse_amount("663,888", Grouping.ENGLISH) == Decimal(663888)
    assert parse_amount("663,888", Grouping.RUSSIAN) == Decimal("663.888")


def test_brackets_mean_minus() -> None:
    """Скобки — знак минус, и он остаётся при величине.

    В отчётности по МСФО состав итога печатается со знаком: «Выручка 89 187,
    Себестоимость (76 881), Валовая прибыль 12 306». Хранить величину расхода
    без знака, как в РСБУ, здесь нельзя — знак несёт сама отчётность, и
    отбросив его, мы получали валовую прибыль 166 068 вместо 12 306.
    """
    assert parse_amount("(29 390)", Grouping.RUSSIAN) == Decimal(-29390)
    assert parse_amount("(29,390)", Grouping.ENGLISH) == Decimal(-29390)


def test_non_numeric_token_is_not_a_number() -> None:
    """Прочерк и текст числами не становятся."""
    for token in ("—", "-", "", "н/д", "1 2 3,4,5"):
        assert parse_amount(token, Grouping.RUSSIAN) is None


def test_minus_sign_survives_parsing() -> None:
    """Математический минус читается наравне с дефисом."""
    assert parse_amount("−15 008", Grouping.RUSSIAN) == Decimal(-15008)
    assert parse_amount("-15,008", Grouping.ENGLISH) == Decimal(-15008)


# --- правдоподобие выбранной конвенции ----------------------------------------


def test_plausibility_confirms_consistent_totals() -> None:
    """Согласованные величины подтверждают конвенцию и называют их число."""
    found = check_plausibility(
        {
            "ifrs.total_non_current_assets": Decimal(400_000),
            "ifrs.total_current_assets": Decimal(263_888),
            "ifrs.total_assets": Decimal(663_888),
            "ifrs.total_equity_and_liabilities": Decimal(663_888),
        },
        revenue=Decimal(507_718),
    )
    assert found.plausible
    assert found.checked == 3
    assert "сверено величин: 3" in found.describe()


def test_plausibility_catches_a_thousandfold_split() -> None:
    """Часть чисел, прочитанная по другой конвенции, расходится на три порядка.

    Ровно то, что случилось бы с ФосАгро при русском прочтении: «663,888»
    стало бы 663,888, а «1 234 567» осталось бы собой.
    """
    found = check_plausibility(
        {
            "ifrs.total_non_current_assets": Decimal(400_000),
            "ifrs.total_current_assets": Decimal(263_888),
            "ifrs.total_assets": Decimal("663.888"),
        }
    )
    assert not found.plausible
    assert found.checked == 1
    assert "кратно тысяче" in found.problems[0]


def test_plausibility_catches_implausible_assets_to_revenue() -> None:
    """Валюта баланса против порядка выручки: расхождение видно сразу."""
    found = check_plausibility(
        {"ifrs.total_assets": Decimal(663_888)}, revenue=Decimal("507.718")
    )
    assert not found.plausible
    assert "границ правдоподобия" in found.problems[0]


def test_plausibility_says_when_it_checked_nothing() -> None:
    """Проверять было нечего — так и сказано, а не «подтверждено»."""
    found = check_plausibility({})
    assert found.checked == 0
    assert "не проверялось" in found.describe()


def test_appendix_prints_the_convention(db_conn, tmp_path) -> None:
    """Конвенция печатается в приложении наравне с единицей измерения.

    Читатель должен иметь возможность проверить, как прочитаны числа: ошибка
    в разделителе разрядов даёт верные коэффициенты при неверных величинах,
    и по самому документу её иначе не увидеть.
    """
    from docx import Document

    from finlib.report.appendix import GROUPING_NAMES, GROUPING_NOT_APPLICABLE
    from finlib.report.document import build_report

    report = build_report("2100010824", db_conn, directory=tmp_path, with_text=False)
    text = "\n".join(item.text for item in Document(report.path).paragraphs)
    assert "Запись чисел в исходном документе" in text
    # Комплект получен из ГИР БО: числа приходят машиночитаемо, и графа
    # объясняет, почему конвенции нет, вместо того чтобы пустовать.
    assert GROUPING_NOT_APPLICABLE in text
    assert not any(name in text for name in GROUPING_NAMES.values())


def test_stored_convention_reaches_the_appendix(db_conn, tmp_path) -> None:
    """Определённая конвенция доходит из src_file до документа."""
    from docx import Document

    from finlib.db import execute
    from finlib.report.appendix import GROUPING_NAMES
    from finlib.report.document import build_report

    execute(
        "UPDATE src_file SET digit_grouping = 'english' "
        "WHERE inn = '2100010824' AND is_actual",
        conn=db_conn,
    )
    report = build_report("2100010824", db_conn, directory=tmp_path, with_text=False)
    text = "\n".join(item.text for item in Document(report.path).paragraphs)
    assert GROUPING_NAMES["english"] in text


def test_ordinary_mismatch_is_not_a_grouping_problem() -> None:
    """Несошедшийся итог сам по себе о конвенции не говорит.

    Проверяется именно кратность тысяче: расхождение в полтора раза бывает
    дефектом отчётности и ловится контролем сходимости, а не этим.
    """
    found = check_plausibility(
        {
            "ifrs.total_non_current_assets": Decimal(400_000),
            "ifrs.total_current_assets": Decimal(263_888),
            "ifrs.total_assets": Decimal(900_000),
        }
    )
    assert found.plausible
