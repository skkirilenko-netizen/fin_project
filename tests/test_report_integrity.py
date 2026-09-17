"""Сквозная сверка чисел: проверенный текст против готового документа.

Дефект, ради которого проверка введена: очистка разметки съедала минус
величины, и «(-0,41)» превращалось в «(0,41)». Число меняло знак уже после
того, как постпроверка его подтвердила, и документ выглядел проверенным.
"""

from datetime import date
from decimal import Decimal

import pytest
from docx import Document

from finlib.llm.service import Conclusion
from finlib.report.document import build_report
from finlib.report.integrity import (
    NumbersAlteredError,
    check_numbers,
    compare,
    numbers_of,
)

FULL_INN = "7736050003"

ANSWER = """### 2. Фактическая база

Валюта баланса (1600) — 25 736 328 136 тыс. руб.

### 3. Аналитическая интерпретация

Ликвидность (cur_liq) снизилась с 1,23 до 0,82 (cur_liq_chg_abs -0,41).

### 4. Риски

Чистый оборотный капитал (nwc) отрицателен: -521 415 920 тыс. руб.

### 5. Ограничения анализа

- Отраслевой привязки нет.

### 6. Вопросы к организации

1. Чем вызвано снижение ликвидности?
"""


# --- разбор чисел ------------------------------------------------------------


def test_numbers_are_read_with_their_signs() -> None:
    """Знак — часть величины, а не оформление."""
    found = numbers_of("снизилась на -0,41 и выросла на 0,41")
    assert found == [Decimal("-0.41"), Decimal("0.41")]


def test_russian_formatting_is_understood() -> None:
    """Разряды пробелами и запятая как десятичный знак разбираются."""
    assert numbers_of("25 736 328 136 тыс. руб.") == [Decimal("25736328136")]
    assert numbers_of("25 736 328 136") == [Decimal("25736328136")]


# --- сверка ------------------------------------------------------------------


def test_identical_text_passes() -> None:
    """Неизменённый текст расхождений не даёт."""
    assert compare("снизилась на -0,41", "снизилась на -0,41") == []


def test_cleanup_removing_codes_passes() -> None:
    """Снятие кодов числа не меняет: документ беднее исходника, и это норма."""
    verified = "Ликвидность (cur_liq) снизилась на -0,41 (cur_liq_chg_abs -0,41)."
    rendered = "Ликвидность снизилась на -0,41 (-0,41)."
    assert compare(verified, rendered) == []


def test_sign_change_is_detected() -> None:
    """Смена знака — отдельный вид расхождения, а не просто чужое число."""
    problems = compare("снизилась на -0,41", "снизилась на 0,41")
    assert [item.kind for item in problems] == ["sign_changed"]
    assert "сменило знак" in problems[0].message


def test_added_number_is_detected() -> None:
    """Число, появившееся после проверки, тоже расхождение."""
    problems = compare("валюта баланса 100", "валюта баланса 100 и прибыль 55")
    assert [item.kind for item in problems] == ["number_added"]


def test_repeated_number_must_not_multiply() -> None:
    """Сверяются мультимножества: одно число не превращается в два."""
    assert compare("0,41", "0,41 и 0,41")


def test_reordering_is_allowed() -> None:
    """Перестановка абзацев оформлением числа не меняет."""
    assert compare("сначала 1,5 потом 2,5", "сначала 2,5 потом 1,5") == []


def test_check_raises_with_sign_change_first() -> None:
    """Смена знака выносится вперёд: это ложное утверждение, а не пропуск."""
    with pytest.raises(NumbersAlteredError) as info:
        check_numbers("было -0,41 и 100", "стало 0,41 и 100 и 7")
    assert "сменило знак" in info.value.problems[0]
    assert len(info.value.problems) == 2


# --- документ ----------------------------------------------------------------


def conclusion_for(text: str) -> Conclusion:
    """Заключение с проверенным текстом и очищенным результатом."""
    from finlib.llm.cleanup import strip_identifiers

    return Conclusion(
        inn=FULL_INN,
        report_date=date(2025, 12, 31),
        text=strip_identifiers(text),
        model="тестовая-модель",
        attempt=1,
        checked_numbers=4,
        verified_text=text,
    )


def test_document_with_intact_numbers_is_written(db_conn, tmp_path) -> None:
    """Исправная сборка проходит сквозную сверку."""
    report = build_report(
        FULL_INN, db_conn, conclusion=conclusion_for(ANSWER), directory=tmp_path
    )
    assert report.path.exists()
    text = "\n".join(item.text for item in Document(report.path).paragraphs)
    assert "-0,41" in text, "величина со знаком дошла до документа"


def test_broken_cleanup_blocks_the_document(db_conn, tmp_path, monkeypatch) -> None:
    """Очистка, портящая знак, документ собрать не даёт.

    Ровно тот дефект, ради которого сверка введена: «(-0,41)» превращалось
    в «(0,41)» уже после постпроверки.
    """
    import finlib.llm.cleanup as cleanup

    original = cleanup.strip_identifiers

    def breaking(text: str, codes=None) -> str:
        """Снимает коды и заодно теряет минус — прежнее поведение очистки."""
        return original(text, codes).replace("(-0,41)", "(0,41)")

    conclusion = Conclusion(
        inn=FULL_INN,
        report_date=date(2025, 12, 31),
        text=breaking(ANSWER),
        model="тестовая-модель",
        attempt=1,
        checked_numbers=4,
        verified_text=ANSWER,
    )
    monkeypatch.setattr(cleanup, "strip_identifiers", breaking)

    with pytest.raises(NumbersAlteredError) as info:
        build_report(FULL_INN, db_conn, conclusion=conclusion, directory=tmp_path)
    assert any("сменило знак" in item for item in info.value.problems)
    assert not list(tmp_path.glob("*.docx")), "испорченный документ остался на диске"


def test_broken_formatting_blocks_the_document(db_conn, tmp_path, monkeypatch) -> None:
    """Искажение на любом шаге после проверки, не только в очистке."""
    import finlib.report.document as module

    original = module._write_sections

    def breaking(document, sections, data):
        """Оформление, подменяющее величину.

        Портится число раздела 3: разделы 2, 4 и 6 собирает расчёт, их
        величины в сверку не входят — сверяется написанное моделью.
        """
        damaged = [
            type(item)(
                item.number,
                item.title,
                tuple(text.replace("0,82", "0,83") for text in item.paragraphs),
            )
            for item in sections
        ]
        original(document, damaged, data)

    monkeypatch.setattr(module, "_write_sections", breaking)

    with pytest.raises(NumbersAlteredError) as info:
        build_report(
            FULL_INN, db_conn, conclusion=conclusion_for(ANSWER), directory=tmp_path
        )
    assert any("0.83" in item.replace(" ", "") for item in info.value.problems)


def test_appendix_numbers_are_not_compared(db_conn, tmp_path) -> None:
    """Сверяются разделы модели, а не расчётная часть.

    В приложении сотни величин, которых в тексте модели нет и быть
    не должно, — сверять их с ним бессмысленно.
    """
    report = build_report(
        FULL_INN, db_conn, conclusion=conclusion_for(ANSWER), directory=tmp_path
    )
    text = "\n".join(item.text for item in Document(report.path).paragraphs)
    assert "Версия методики оценки" in text, "приложение на месте и сверку не сорвало"


def test_missing_sections_block_the_document(db_conn, tmp_path, monkeypatch) -> None:
    """Пустая сверка не вправе выглядеть успехом.

    Разбор документа на разделы держится на формате заголовка. Изменится
    формат — функция вернёт пустую строку, сверять станет нечего, и документ
    выйдет с пометкой «числа проверены»: ровно та конструкция, из-за которой
    контроль утверждений текста месяцами не выполнялся. Поэтому число
    найденных разделов сверяется с ожидаемым.
    """
    import finlib.report.document as module

    # Заголовки в документе те же, а ожидание разделов разошлось с ними:
    # так же выглядит и обратный случай — изменившийся формат заголовка
    # при прежнем ожидании.
    monkeypatch.setattr(module, "EXPECTED", (*module.EXPECTED, (9, "Небывалый")))

    with pytest.raises(module.SectionsNotFoundError) as info:
        build_report(
            FULL_INN, db_conn, conclusion=conclusion_for(ANSWER), directory=tmp_path
        )
    assert "9" in str(info.value)
    assert not list(tmp_path.glob("*.docx")), "документ не остаётся на диске"


def test_section_scan_reports_what_it_did_not_find(tmp_path) -> None:
    """Разбор называет, каких разделов не нашёл, а не возвращает пустоту."""
    from finlib.report.document import SectionsNotFoundError, _model_text_of

    document = Document()
    document.add_heading("Раздел 3: Аналитическая интерпретация", level=1)
    document.add_paragraph("Коэффициент текущей ликвидности 1,49.")
    path = tmp_path / "renamed.docx"
    document.save(path)

    with pytest.raises(SectionsNotFoundError) as info:
        _model_text_of(path)
    assert "3" in str(info.value) and "5" in str(info.value)
