"""Критерий выхода «эмитенты закрыты»: пять критериев на дату — факт и выполнен ли.

    uv run python eval/exit_criteria.py --as-of ГГГГ-ММ-ДД \
        [--answers КАТАЛОГ] [--doc ПУТЬ] [--out ПУТЬ.md]

Критерий выхода из части первой (`ROADMAP.md`, «Критерий выхода — „эмитенты
закрыты“») — пять условий, и переход к части второй — когда выполнены все:

- **охват** — каждый эмитент с выпусками в обращении в круге, с корзиной
  или статусом: перечень Cbonds в обращении против точек маршрута дня
  (`routing_day`) и журнала исключений (`routing_store.exclusions`);
- **качество** — на календаре событий за двенадцать месяцев до даты ни одного
  дефолта в «Без внимания» накануне события: календарь —
  `market_lead_run.events` (тот же, что у замеров), корзина накануне —
  правило замера фазы 5 (`ltm_basis_run.eve_basket`), отдельно по пересчёту
  (`backfill`) и по наблюдению (`routing_day`); «Без внимания» хоть в одном
  роде — не выполнен (решение владельца 12.10.2026);
- **устойчивость** — чистые прогоны по расписанию подряд не меньше месяца:
  счёт — `status_run._streak`, тот же, что у утренней сводки. День чист,
  если чист прогон по расписанию этого дня: ручной повтор после упавшего
  или незавершённого планового день не спасает, повтор после успешного
  и автоповтор стадии внутри прогона (run-retry, `retries` в записи
  доставки) счёт не обрывают (решение владельца 12.10.2026);
- **экспертиза** — последние проверки подряд пройдены: файлы ответов
  (пакеты `eval/expert_review_run.py` с отметками эксперта); условие
  и с какого пакета проверки идут в счёт — `methodology/acceptance.yaml`;
- **документация** — обзор методики для аналитика называет каждое основание
  `routing.yaml` кодом и не помечен черновиком.

**Замер не считает сам.** Каждый критерий берёт величину у того кода,
который её считает для системы; здесь — только сведение и сравнение
с условием. **Только чтение**: база открывается на чтение, в неё не пишется
ничего, сеть не нужна — перечня на диске нет, и критерий печатает «не
установлено», а не идёт за ним в Cbonds.

**«Не установлено» — третий исход, а не «нет».** Критерий, для которого
нет данных (нет файлов ответов, нет перечня на диске, нет событий с историей),
не выполнен, но и не провален: печатается, чего не хватает.

**Ограничения названы в выводе**: перечень в обращении, карточки и календарь
событий — с диска на день запуска, а не на дату: истории перечня нет.
Часть критерия качества «выявляемость, прирост и упреждение измерены
и записаны» — по документам (`ROADMAP.md`, фазы 3 и 6), скрипт её
не проверяет.
"""

import argparse
import contextlib
import logging
import re
import sys
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from change_report_run import _read  # noqa: E402
from ltm_basis_run import eve_basket  # noqa: E402
from status_run import _streak  # noqa: E402

from finlib.acceptance import ExpertReviewAcceptance, load_acceptance  # noqa: E402
from finlib.config import settings  # noqa: E402
from finlib.db import PgConnection, connection, fetch_all  # noqa: E402

logger = logging.getLogger(__name__)

# Окна критериев — слова дорожной карты («за последние двенадцать месяцев»,
# «не меньше месяца»), а не пороги маршрута: в методику они не вносятся.
QUALITY_MONTHS = 12
STABILITY_MONTHS = 1

# **Слова ответа эксперта — те же, что у пакета** (`eval/expert_review_run.py`,
# `MARKS`, `CRITICAL_YES`, `SIGNS`). Пакет и этот разбор пришли разными
# ветками; после слияния обеих слова берутся импортом оттуда, а не
# повторяются здесь. Условие проверки (доля «лишних», критичных, подряд,
# с какого пакета) — в методике, `methodology/acceptance.yaml`. Перечень
# допустимых слов читается из шапки самой графы вердикта — у старого пакета
# он свой, и слово вне перечня делает проверку незавершённой, а не теряется.
VERDICT_HEAD = "вердикт эксперта"
CRITICAL_HEAD = "критичный дефект"
EXTRA = "лишний"
CRITICAL_YES = "да"
SIGN_CLEAR_SEEN = "«Без внимания» просмотрен целиком (да / нет)"
REVIEW, CLEAR, SAMPLE = "Разбор", "Без внимания", "Выборка"
DRAFT = "Черновик"

YES, NO, UNKNOWN = "да", "нет", "не установлено"


@dataclass(frozen=True, slots=True)
class Criterion:
    """Строка итога: критерий, факт и исход (None — не установлено)."""

    name: str
    fact: str
    done: bool | None


def months_back(day: date, months: int) -> date:
    """Тот же день на `months` месяцев раньше; дня нет в месяце — последний день месяца."""
    month = day.month - months
    year = day.year + (month - 1) // 12
    month = (month - 1) % 12 + 1
    for last in (31, 30, 29, 28):
        try:
            return date(year, month, min(day.day, last))
        except ValueError:
            continue
    raise ValueError(f"нет даты {year}-{month}")  # pragma: no cover


def _ru(day: date) -> str:
    """Дата так, как её печатают отчёты: «13.10.2026»."""
    return f"{day:%d.%m.%Y}"


def coverage(
    conn: PgConnection,
    as_of: date,
    universe: Mapping[str, str] | None,
    excluded: Iterable[str],
) -> Criterion:
    """Охват: каждый эмитент перечня в обращении — с точкой маршрута дня либо вышел записью."""
    name = "охват"
    if universe is None:
        return Criterion(name, "перечня выпусков в обращении на диске нет", None)
    points = _read(conn, "run", as_of)
    if not points:
        return Criterion(name, f"маршрута на {_ru(as_of)} в истории нет (routing_day)", None)
    excluded = set(excluded)
    routed = [inn for inn in universe if inn in points]
    left = [inn for inn in universe if inn not in points and inn in excluded]
    missing = sorted(inn for inn in universe if inn not in points and inn not in excluded)
    status = sum(1 for inn in routed if points[inn]["basket"] == "status_unknown")
    fact = (
        f"с выпусками в обращении {len(universe)}; с корзиной на {_ru(as_of)} "
        f"{len(routed)} (из них «Установить статус» {status}); вышли записью "
        f"в журнале исключений {len(left)}; без корзины {len(missing)}"
    )
    if missing:
        shown = ", ".join(missing[:10]) + (" …" if len(missing) > 10 else "")
        fact += f": {shown}"
    return Criterion(name, fact, not missing)


_FIRST = "SELECT min(as_of) AS first FROM {source} WHERE TRUE {kind}"
_SERIES = "SELECT as_of, basket FROM {source} WHERE inn = %(inn)s {kind} ORDER BY as_of"
_SOURCES = {
    "пересчёт": ("routing_history", "AND kind = 'backfill'"),
    "наблюдение": ("routing_day", ""),
}


def quality(conn: PgConnection, as_of: date, happened: Mapping[str, date] | None) -> Criterion:
    """Качество: у эмитента с событием за двенадцать месяцев накануне — не «Без внимания»."""
    name = "качество"
    if happened is None:
        return Criterion(name, "календаря событий на диске нет", None)
    since = months_back(as_of, QUALITY_MONTHS)
    window = {inn: moment for inn, moment in happened.items() if since < moment <= as_of}
    seen = dict.fromkeys(_SOURCES, 0)
    missed: dict[str, list[str]] = {source: [] for source in _SOURCES}
    for source, (table, kind) in _SOURCES.items():
        [row] = fetch_all(_FIRST.format(source=table, kind=kind), {}, conn=conn)
        first = row["first"]
        for inn, moment in sorted(window.items()):
            points = [
                (point["as_of"], point["basket"])
                for point in fetch_all(
                    _SERIES.format(source=table, kind=kind), {"inn": inn}, conn=conn
                )
            ]
            # Правило замера фазы 5 — его же кодом (`ltm_basis_run.eve_basket`).
            basket = eve_basket(points, moment, first) if points else None
            if basket is None:
                continue
            seen[source] += 1
            if basket == "clear":
                missed[source].append(f"{inn} ({_ru(moment)})")
    fact = (
        f"событий с {_ru(since + timedelta(days=1))} по {_ru(as_of)}: {len(window)}; "
        + "; ".join(
            f"{source}: с историей до события {seen[source]}, «Без внимания» "
            f"накануне {len(missed[source])}"
            + (f" — {', '.join(missed[source])}" if missed[source] else "")
            for source in _SOURCES
        )
        + ". Выявляемость, прирост и упреждение — по документам (ROADMAP, фазы 3 и 6)"
    )
    if not any(seen.values()):
        return Criterion(name, fact, None)
    return Criterion(name, fact, not any(missed.values()))


_RUNS = """
SELECT id, kind, as_of, status, started_at, finished_at, sources, note
FROM routing_run WHERE started_at < %(until)s ORDER BY id DESC
"""


def _working_days(since: date, until: date) -> int:
    """Рабочих дней (пн–пт) в полуинтервале (since, until]."""
    return sum(
        1
        for offset in range(1, (until - since).days + 1)
        if (since + timedelta(days=offset)).weekday() < 5
    )


def stability(conn: PgConnection, as_of: date) -> Criterion:
    """Устойчивость: чистые прогоны по расписанию подряд — не меньше месяца до даты."""
    name = "устойчивость"
    end = datetime.combine(as_of, time(23, 59, 59)).astimezone()
    runs = fetch_all(_RUNS, {"until": end + timedelta(seconds=1)}, conn=conn)
    count, broken = _streak(runs, end)
    since = months_back(as_of, STABILITY_MONTHS)
    need = _working_days(since, as_of)
    fact = (
        f"чистых прогонов по расписанию подряд на {_ru(as_of)}: {count} рабочих дней "
        f"(нужно {need} — с {_ru(since + timedelta(days=1))}); обрыв счёта: {broken}"
    )
    return Criterion(name, fact, count >= need)


@dataclass(frozen=True, slots=True)
class Review:
    """Одна экспертная проверка по файлу ответа."""

    path: Path
    day: date | None
    complete: bool
    critical: int
    extra: int
    review_rows: int
    why: str

    def passed(self, rule: ExpertReviewAcceptance) -> bool:
        """Пройдена: размечена целиком, критичных не больше допуска, «лишних» не больше доли."""
        if not self.complete or self.critical > rule.critical_max:
            return False
        return Decimal(self.extra) <= rule.extra_share_max * self.review_rows


def _column(head: list, start: str) -> int | None:
    """Номер графы, чья шапка начинается словами; нет такой — None."""
    return next(
        (number for number, cell in enumerate(head) if str(cell or "").startswith(start)),
        None,
    )


def _marks(head: str) -> set[str]:
    """Допустимые слова вердикта из шапки графы: «вердикт эксперта (а / б)»."""
    found = re.search(r"\((.*)\)", head)
    return {word.strip() for word in found.group(1).split("/")} if found else set()


def read_review(path: Path) -> Review:
    """Разбор файла ответа: дата маршрута, разметка листов, критичные и «лишние»."""
    book = load_workbook(path, read_only=True, data_only=True)
    sheets = set(book.sheetnames)
    summary = (
        {row[0]: row[1] for row in book["Сводка"].iter_rows(values_only=True) if row and row[0]}
        if "Сводка" in sheets
        else {}
    )
    day = None
    raw = str(summary.get("Дата маршрута") or "")
    with contextlib.suppress(ValueError):
        day = datetime.strptime(raw, "%d.%m.%Y").date()
    problems: list[str] = []
    if day is None:
        problems.append("нет даты маршрута в «Сводке»")
    critical = extra = review_rows = 0
    for title in (REVIEW, CLEAR, SAMPLE):
        if title not in sheets:
            problems.append(f"нет листа «{title}»")
            continue
        head, *rows = [list(row) for row in book[title].iter_rows(values_only=True)]
        verdict, defect = _column(head, VERDICT_HEAD), _column(head, CRITICAL_HEAD)
        if verdict is None or defect is None:
            problems.append(f"на листе «{title}» нет граф вердикта")
            continue
        allowed = _marks(str(head[verdict]))
        rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
        marked = [str(row[verdict] or "").strip().lower() for row in rows]
        unknown = sorted({mark for mark in marked if mark and mark not in allowed})
        if unknown:
            problems.append(f"«{title}»: слова вне перечня графы — {', '.join(unknown)}")
        empty = sum(1 for mark in marked if not mark)
        if title in (REVIEW, SAMPLE) and empty:
            problems.append(f"«{title}»: не размечено {empty} из {len(rows)}")
        critical += sum(
            1 for row in rows if str(row[defect] or "").strip().lower() == CRITICAL_YES
        )
        if title == REVIEW:
            review_rows = len(rows)
            extra = sum(1 for mark in marked if mark == EXTRA)
    signs = (
        {row[0]: row[1] for row in book["Инструкция"].iter_rows(values_only=True) if row and row[0]}
        if "Инструкция" in sheets
        else {}
    )
    if str(signs.get(SIGN_CLEAR_SEEN) or "").strip().lower() != YES:
        problems.append("не отмечено, что «Без внимания» просмотрен целиком")
    book.close()
    return Review(
        path=path,
        day=day,
        complete=not problems,
        critical=critical,
        extra=extra,
        review_rows=review_rows,
        why="; ".join(problems),
    )


def expertise(answers: Path, rule: ExpertReviewAcceptance) -> Criterion:
    """Экспертиза: последние `rule.consecutive` проверок подряд пройдены; счёт — с `counted_from`.

    **Проверки до `counted_from` в счёт не идут** (решение владельца
    12.10.2026: счёт — с пакета 13.10.2026): они печатаются строкой «не в
    счёт», но ни продолжить серию, ни оборвать её не могут. Проверка без
    даты маршрута в счёт не идёт тоже — её место в серии неизвестно.
    """
    name = "экспертиза"
    files = sorted(answers.glob("*.xlsx")) if answers.is_dir() else []
    files = [path for path in files if not path.name.startswith("~$")]
    if not files:
        return Criterion(name, f"файлов ответов нет ({answers})", None)
    reviews = sorted(
        (read_review(path) for path in files),
        key=lambda item: (item.day or date.min, item.path.name),
    )
    counted = [item for item in reviews if item.day and item.day >= rule.counted_from]
    lines = []
    for item in reviews:
        when = _ru(item.day) if item.day else "дата не установлена"
        share = (
            f"«лишних» в «Разборе» {item.extra} из {item.review_rows}"
            if item.review_rows
            else "«Разбор» пуст"
        )
        if not item.complete:
            verdict = f"не завершена ({item.why})"
        else:
            verdict = "пройдена" if item.passed(rule) else "не пройдена"
        if item not in counted:
            verdict += f"; не в счёт — до {_ru(rule.counted_from)}"
        lines.append(
            f"{when} ({item.path.name}): {verdict}; критичных {item.critical}; {share}"
        )
    last = counted[-rule.consecutive :] if rule.consecutive else []
    done = len(last) == rule.consecutive and all(item.passed(rule) for item in last)
    fact = (
        f"проверок {len(reviews)}, в счёт с {_ru(rule.counted_from)} — {len(counted)}; "
        f"условие — последние {rule.consecutive} подряд пройдены (критичных не больше "
        f"{rule.critical_max}, «лишних» в «Разборе» не больше "
        f"{rule.extra_share_max * 100:.0f} %, размечено целиком; "
        f"methodology/acceptance.yaml): " + " | ".join(lines)
    )
    if not counted:
        return Criterion(name, fact, None)
    return Criterion(name, fact, done)


def _ground_codes(routing: Mapping) -> set[str]:
    """Коды оснований справочника маршрута: корзины, справочные, типы, источники."""
    codes = {
        ground["code"]
        for basket in routing.get("baskets") or []
        for ground in basket.get("grounds") or []
    }
    codes |= {item["code"] for item in routing.get("reference") or []}
    codes |= {kind["ground"] for kind in routing.get("issuer_types") or []}
    codes |= set(routing.get("ground_sources") or {})
    return codes


def documentation(doc: Path, routing_yaml: Path) -> Criterion:
    """Документация: обзор методики называет каждое основание кодом и не помечен черновиком."""
    name = "документация"
    if not doc.exists():
        return Criterion(name, f"обзора методики нет ({doc})", False)
    text = doc.read_text(encoding="utf-8")
    codes = _ground_codes(yaml.safe_load(routing_yaml.read_text(encoding="utf-8")))
    missing = sorted(code for code in codes if f"`{code}`" not in text)
    draft = DRAFT in "\n".join(text.splitlines()[:10])
    fact = (
        f"{doc.name}: оснований routing.yaml названо кодом {len(codes) - len(missing)} "
        f"из {len(codes)}"
        + (f" (нет: {', '.join(missing)})" if missing else "")
        + ("; помечен черновиком — утверждает владелец" if draft else "; черновиком не помечен")
    )
    return Criterion(name, fact, not missing and not draft)


def _disk() -> tuple[Mapping[str, str] | None, set[str], Mapping[str, date] | None]:
    """Перечень в обращении, вышедшие записью и календарь событий — с диска, без сети."""
    from finlib.scoring.routing import load_routing
    from finlib.scoring.routing_store import cards, exclusions
    from finlib.sources import cbonds

    # Перечня на диске нет — критерии «не установлено»: идти за ним в Cbonds
    # замер не вправе (`cbonds.fetch` пошёл бы в сеть при пустом кэше).
    if not cbonds._cache_path("emissions_ru_outstanding").exists():
        return None, set(), None
    from market_lead_run import events

    excluded, _ = exclusions(cards(), load_routing())
    return cbonds.bond_issuers(), set(excluded), events()


def build(
    conn: PgConnection,
    as_of: date,
    answers: Path,
    doc: Path,
    universe: Mapping[str, str] | None,
    excluded: Iterable[str],
    happened: Mapping[str, date] | None,
) -> list[Criterion]:
    """Пять критериев на дату по готовым входам с диска."""
    return [
        coverage(conn, as_of, universe, excluded),
        quality(conn, as_of, happened),
        stability(conn, as_of),
        expertise(answers, load_acceptance()),
        documentation(doc, doc.parent.parent / "methodology" / "routing.yaml"),
    ]


def render(as_of: date, rows: list[Criterion]) -> str:
    """Таблица «критерий / факт / выполнен» с итогом."""
    shown = {True: YES, False: NO, None: UNKNOWN}
    lines = [
        f"# Критерий выхода «эмитенты закрыты» на {_ru(as_of)}",
        "",
        "| Критерий | Факт | Выполнен |",
        "|---|---|---|",
        *[
            f"| {row.name} | {row.fact.replace('|', '/')} | {shown[row.done]} |"
            for row in rows
        ],
        "",
        f"**Выполнено {sum(1 for row in rows if row.done)} из {len(rows)}**; переход "
        "к части второй — когда выполнены все пять (ROADMAP).",
        "",
        "Ограничения: перечень в обращении, карточки и календарь событий — с диска "
        "на день запуска, а не на дату; «не установлено» — данных для критерия нет, "
        "это не «выполнен» и не «провален».",
        "",
    ]
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    """Печатает таблицу критериев; с `--out` — пишет её в файл, готовый не переписывает."""
    logging.basicConfig(level=logging.ERROR, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--as-of", type=date.fromisoformat, required=True)
    parser.add_argument(
        "--answers", type=Path, default=settings.data_dir / "output" / "expert_answers"
    )
    parser.add_argument(
        "--doc", type=Path, default=settings.base_dir / "docs" / "methodology_overview.md"
    )
    parser.add_argument("--out", type=Path)
    args = parser.parse_args(argv)
    if args.out is not None and args.out.exists():
        print(f"итог {args.out} уже есть и не переписывается: выберите другой путь")
        return 1
    universe, excluded, happened = _disk()
    with connection() as conn:
        rows = build(conn, args.as_of, args.answers, args.doc, universe, excluded, happened)
        conn.rollback()
    text = render(args.as_of, rows)
    if args.out is not None:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(text, encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
