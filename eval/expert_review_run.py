"""Пакет для экспертной проверки списка: «Разбор», «Без внимания» и выборка.

    uv run python eval/expert_review_run.py --as-of ГГГГ-ММ-ДД --out ПУТЬ.xlsx

Критерий выхода ROADMAP — «экспертная проверка не находит критичных
дефектов два раза подряд»; пакет — материал одной такой проверки.

**Пакет воспроизводим** (08.10.2026): маршрут дня читается из истории
(`routing_day` — последняя точка дня, та же, что у отчёта изменений),
а не пересчитывается на сегодня. Пакет 30.09.2026 собирался `routing_rows`
на дату запуска, и пересобрать его тем же видом позже было нечем: маршрут
к тому дню был уже другим. Теперь та же дата и та же база дают тот же
пакет. **Только чтение**: в базу не пишется ничего. Дня в истории нет —
отказ, а не пустой пакет; готовый файл не переписывается.

Листы:

- «Сводка» — дата, числа отбора, знаменатели, ограничения;
- «Разбор» — все эмитенты корзины;
- «Без внимания» — то же, на полный просмотр (решение владельца 30.09.2026);
- «Выборка» — 30 из «Без внимания» на углублённую проверку, с основанием
  отбора.

На каждом листе — корзина, решающее основание, все основания, величины
решения с единицами (печать справочником своего стандарта — той же точкой,
что у списка), ссылка на карточку (`cards/<ИНН>.md` рядом с файлом пакета,
если она там лежит) и пустые графы вердикта эксперта.

**Состав выборки — решение владельца 30.09.2026.** Сперва все эмитенты
«Без внимания» с крупным долгом по облигациям — тем же, что у маршрута:
верхняя доля по объёму облигаций в обращении (`routing.yaml`,
`systemic.top_share`), тем же кодом (`routing_store._outstanding`,
`_top_share`). Затем до тридцати — эмитенты с наибольшим долгом
по отчётности: совокупный долг базы (`debt_total` из величин решения
точки), приведённый к рублям по единице комплекта. Отбор детерминирован:
добор — по убыванию долга, ничья — по ИНН. Это отбор для проверки,
а не правило маршрута, и в методику он не вносится.

**Ограничение воспроизводимости названо в сводке**: объём облигаций
в обращении и наименования берутся с диска на день сборки, а не на дату
маршрута, — истории объёмов нет. Крупный долг по облигациям у пакета
за прошлую дату поэтому может разойтись с тем, что видел маршрут того дня.

**Выпусков в обращении по источникам** (12.10.2026): графа «файл / перечень
/ срез» — число выпусков в обращении у файла выпусков эмитента, перечня
Cbonds и последнего среза биржи не позже дня маршрута, и, если есть,
по поручительству за чужие выпуски. Считает тот же код, что периметр
(`routing_store.bond_sources`): эксперт, отмечающий «нет выпусков», видит
то же, что увидел бы маршрут, а расхождение источников видно строкой.
Файл и перечень — с диска на день сборки, как объём.

**Лист «Инструкция»** (12.10.2026) — для коллег: что проверить на каждом
листе, как отмечать «лишний» и «нет выпусков», критерий проверки и графы
отметки эксперта. Слова вердикта и критерий объявлены здесь одним местом
(`MARKS`, `EXTRA_SHARE_MAX`); ответы по ним читает `eval/exit_criteria.py`.
"""

import argparse
import logging
import sys
from collections.abc import Mapping
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))

from change_report_run import _ground_names, _read  # noqa: E402
from watchlist_csv import VALUES, _sum  # noqa: E402

from finlib.db import PgConnection, connection  # noqa: E402
from finlib.normalize.lines import load_lines  # noqa: E402
from finlib.scoring.routing import load_routing  # noqa: E402
from finlib.scoring.routing_catalogue import catalogue_for  # noqa: E402
from finlib.scoring.routing_store import (  # noqa: E402
    BACKING,
    BOARD,
    LISTED,
    OWN_FILE,
    _outstanding,
    _top_share,
    bond_sources,
    cards,
)
from finlib.sources.cbonds_events import OKEI_MULTIPLIER  # noqa: E402
from finlib.sources.market import universe  # noqa: E402
from finlib.standards import Standard  # noqa: E402

logger = logging.getLogger(__name__)

SAMPLE_SIZE = 30
# Долг по отчётности — величина решения своего стандарта: совокупные
# заёмные средства базы (LTM на промежуточную дату либо годовая).
DEBT = "debt_total"
BASKETS = (("review", "Разбор"), ("clear", "Без внимания"))

HEADER = (
    "ИНН",
    "наименование",
    "корзина",
    "решающее основание",
    "все основания корзины",
    "прочие сработавшие основания",
    "подгруппа",
    "действие",
    "отчётная дата",
    "стандарт",
    "единица",
    "крупный долг (верхняя доля по облигациям)",
    "объём облигаций в обращении, руб.",
    "выпусков в обращении: файл / перечень / срез",
    *VALUES,
    "денежные средства",
    "платежи 12 месяцев",
    "оферты 12 месяцев",
    "все величины решения",
    "карточка",
)
# **Слова вердикта — одно место.** Их пишет эксперт, а читает разбор
# ответов: написание, разошедшееся между инструкцией и разбором, превратило
# бы «лишний» в нераспознанное и молча выбросило бы его из доли.
AGREE = "согласен"
DISAGREE = "не согласен"
EXTRA = "лишний"
NO_BONDS = "нет выпусков"
MARKS = (AGREE, DISAGREE, EXTRA, NO_BONDS)
CRITICAL_YES = "да"
VERDICT = (
    f"вердикт эксперта ({' / '.join(MARKS)})",
    "критичный дефект (да / нет)",
    "комментарий",
)
# Критерий проверки — решение владельца 12.10.2026: ни одного критичного
# дефекта, «лишних» в «Разборе» не больше десятой части листа. Это критерий
# экспертной проверки, а не порог маршрута, и в методику он не вносится.
EXTRA_SHARE_MAX = Decimal("0.10")
# Отметка эксперта в конце «Инструкции»: подписи в графе A, ответ — в графе B.
SIGN_EXPERT = "Эксперт"
SIGN_DATE = "Дата проверки"
SIGN_CLEAR_SEEN = "«Без внимания» просмотрен целиком (да / нет)"
SIGNS = (SIGN_EXPERT, SIGN_DATE, SIGN_CLEAR_SEEN)
INSTRUCTION = "Инструкция"


class NoRouteError(ValueError):
    """Маршрута на дату в истории нет: пакет собирать не из чего."""


def _rub(value: Decimal | None, unit: str) -> Decimal | None:
    """Денежная величина комплекта в рублях; единица неизвестна — None."""
    if value is None:
        return None
    names = load_lines().units.names
    multiplier = next(
        (factor for code, factor in OKEI_MULTIPLIER.items() if names.get(code) == unit),
        None,
    )
    return None if multiplier is None else value * multiplier


def _inputs(point: Mapping) -> dict:
    """Величины решения точки; пусто — точка записана без них."""
    return point.get("inputs") or {}


def _metrics(point: Mapping) -> dict[str, Decimal]:
    """Показатели решения точки: код → `Decimal`; хранятся строками."""
    raw = _inputs(point).get("metrics") or {}
    return {code: Decimal(value) for code, value in raw.items() if value not in (None, "")}


def _shown(point: Mapping) -> dict[str, tuple[str, str]]:
    """Показатели так, как их печатает справочник стандарта: код → (имя, текст)."""
    if not point.get("standard"):
        return {}
    catalogue = catalogue_for(Standard(point["standard"]))
    unit = str(_inputs(point).get("unit") or "")
    return {
        code: (catalogue.name_of(code), catalogue.shown(code, value, unit))
        for code, value in _metrics(point).items()
    }


def _money(raw: object, unit: str) -> str:
    """Денежная величина решения с единицей; пусто — не раскрыта."""
    return _sum(Decimal(str(raw)), unit) if raw not in (None, "") else ""


def _bonds(said: Mapping[str, int] | None) -> str:
    """Выпуски в обращении по источникам: «файл / перечень / срез», поручительство — если есть."""
    if said is None:
        return ""
    shown = " / ".join(str(said.get(name, 0)) for name in (OWN_FILE, LISTED, BOARD))
    backing = said.get(BACKING, 0)
    return f"{shown}, {BACKING} {backing}" if backing else shown


def _row(
    routing,  # noqa: ANN001
    point: Mapping,
    name: str,
    volume: Decimal | None,
    systemic: bool,
    bonds: Mapping[str, int] | None = None,
) -> list:
    """Строка листа по точке истории."""
    names = _ground_names(routing)
    inputs = _inputs(point)
    unit = str(inputs.get("unit") or "")
    grounds = list(point.get("grounds") or [])
    others = [code for code in point.get("grounds_all") or [] if code not in grounds]
    shown = _shown(point)
    money = inputs.get("refinance") or {}
    money_unit = str(money.get("unit") or unit)
    action = inputs.get("action") or {}
    return [
        point["inn"],
        name,
        routing.basket(point["basket"]).name,
        names.get(grounds[0], grounds[0]) if grounds else "",
        "; ".join(names.get(code, code) for code in grounds),
        "; ".join(names.get(code, code) for code in others),
        str(action.get("subgroup_name") or point.get("subgroup") or ""),
        str(action.get("text") or ""),
        f"{point['report_date']:%Y-%m-%d}" if point.get("report_date") else "",
        point.get("standard") or "",
        unit,
        "да" if systemic else "",
        f"{volume:,.0f}".replace(",", " ") if volume is not None else "",
        _bonds(bonds),
        *[shown[code][1] if code in shown else "" for code in VALUES],
        _money(inputs.get("cash"), unit),
        _money(money.get("due"), money_unit),
        _money(money.get("offered"), money_unit),
        " | ".join(f"{label}: {text}" for label, text in shown.values()),
        "",
    ]


def _sheet(
    book: Workbook, title: str, rows: list[tuple[str, list]], cards_dir: Path,
    extra: tuple[str, ...] = (),
) -> None:
    """Лист с шапкой, ИНН текстом, ссылкой на карточку и пустыми графами вердикта."""
    sheet = book.create_sheet(title)
    columns = [*HEADER, *extra, *VERDICT]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    link_col = len(HEADER)
    for number, (inn, values) in enumerate(rows, start=2):
        sheet.append([*values, *[""] * (len(columns) - len(values))])
        # ИНН текстом: ведущий ноль у четверти организаций.
        sheet.cell(row=number, column=1).number_format = "@"
        if (cards_dir / f"{inn}.md").exists():
            cell = sheet.cell(row=number, column=link_col)
            cell.value = f"cards/{inn}.md"
            cell.hyperlink = f"cards/{inn}.md"
            cell.font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "C2"


def _instruction(book: Workbook, as_of: date, review: int, clear: int, sample: int) -> None:
    """Лист «Инструкция» первым: что проверить, как отмечать, критерий, отметка эксперта."""
    verdict, critical, comment = VERDICT
    limit = f"{EXTRA_SHARE_MAX * 100:.0f} %"
    text: list[tuple[str, str]] = [
        ("Экспертная проверка списка наблюдения", f"маршрут на {as_of:%d.%m.%Y}"),
        ("Зачем", "Критерий выхода «экспертиза»: проверка не находит критичных "
                  "дефектов два раза подряд. Маршрут отвечает на один вопрос — "
                  "нужен ли по эмитенту человек; класс он не присваивает."),
        ("", ""),
        ("Что проверить", ""),
        (f"«Разбор» ({review})", f"Каждую строку: нужен ли по эмитенту человек. "
                                 f"«{AGREE}» — место в «Разборе»; «{EXTRA}» — разбор "
                                 f"не нужен, по вашему знанию эмитента его место ниже "
                                 f"(в комментарии — какая корзина и почему); "
                                 f"«{DISAGREE}» — основание названо неверно (ошибка данных, "
                                 f"не то событие); «{NO_BONDS}» — облигаций в обращении нет."),
        (f"«Без внимания» ({clear})", f"Просмотреть целиком. Отмечать только тех, о ком "
                                      f"знаете проблему: «{DISAGREE}» — человек нужен "
                                      f"(в комментарии — что известно); «{NO_BONDS}» — "
                                      f"облигаций нет. Пустой вердикт — просмотрен, "
                                      f"возражений нет."),
        (f"«Выборка» ({sample})", f"Углублённо, каждую строку: открыть карточку (графа "
                                  f"«карточка»), сверить одну-две величины решения "
                                  f"с отчётностью и основания с тем, что известно "
                                  f"об эмитенте; «{AGREE}», «{DISAGREE}» или «{NO_BONDS}»."),
        ("", ""),
        ("Как отмечать", ""),
        (verdict, f"Одно из слов: {', '.join(f'«{mark}»' for mark in MARKS)} — "
                  "без других слов в ячейке."),
        (f"«{NO_BONDS}»", "Сперва графа «выпусков в обращении: файл / перечень / срез»: "
                          "три числа — выпуски в обращении у файла выпусков эмитента, "
                          "перечня Cbonds и последнего среза биржи. Нули везде — "
                          "маршрут тоже видит «нет выпусков»; числа расходятся — "
                          "отметьте, что знаете вы, и напишите источник в комментарии."),
        (critical, f"«{CRITICAL_YES}» — если по списку человек пропустил бы эмитента "
                   "с проблемой: известный дефолт, неплатёж, банкротство или иное "
                   "кредитное событие у эмитента в «Без внимания» или «Выборке», "
                   "либо ошибка в данных, из-за которой корзина неверна. Иначе «нет» "
                   "или пусто."),
        (comment, f"Обязателен при «{DISAGREE}», «{EXTRA}», «{NO_BONDS}» "
                  f"и критичном дефекте «{CRITICAL_YES}»."),
        ("", ""),
        ("Критерий проверки", f"Пройдена, если: ни одной строки с критичным дефектом "
                              f"«{CRITICAL_YES}»; «{EXTRA}» в «Разборе» — не больше "
                              f"{limit} строк листа; «Разбор» и «Выборка» размечены "
                              f"целиком, «Без внимания» просмотрен целиком. Не размечено "
                              f"целиком — проверка не завершена, а не пройдена."),
        ("Ответ", "Файл не переименовывать и листов не менять; заполнить отметку "
                  "ниже и вернуть файл."),
        ("", ""),
        ("Отметка эксперта", ""),
        *[(sign, "") for sign in SIGNS],
    ]
    sheet = book.create_sheet(INSTRUCTION, 0)
    for label, value in text:
        sheet.append([label, value])
        if label and not value or label == text[0][0]:
            sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
    sheet.column_dimensions["A"].width = 45
    sheet.column_dimensions["B"].width = 110


def build(conn: PgConnection, as_of: date, out: Path) -> list[tuple[str, object]]:
    """Собирает пакет на дату из истории маршрута; возвращает строки сводки."""
    routing = load_routing()
    points = _read(conn, "run", as_of)
    if not points:
        raise NoRouteError(
            f"маршрута на {as_of:%d.%m.%Y} в истории нет (routing_day): "
            "пакет собирать не из чего"
        )
    known = cards()
    named = {inn: str((known.get(inn) or {}).get("name_rus") or "").strip() for inn in points}
    volumes = {
        inn: total
        for inn in universe()
        if (total := _outstanding(inn)) is not None and total > 0
    }
    systemic = _top_share(volumes, routing.systemic.top_share)
    bonds = bond_sources(points, as_of, routing)
    by_basket = {
        code: sorted(
            (point for point in points.values() if point["basket"] == code),
            key=lambda point: (named[point["inn"]], point["inn"]),
        )
        for code, _ in BASKETS
    }
    cards_dir = out.parent / "cards"

    def line(point: Mapping) -> tuple[str, list]:
        inn = point["inn"]
        return inn, _row(
            routing, point, named[inn], volumes.get(inn), inn in systemic, bonds.get(inn)
        )

    book = Workbook()
    book.remove(book.active)
    for code, title in BASKETS:
        _sheet(book, title, [line(point) for point in by_basket[code]], cards_dir)

    clear = by_basket["clear"]
    pool = sorted((p for p in clear if p["inn"] in systemic), key=lambda p: p["inn"])
    others = [p for p in clear if p["inn"] not in systemic]
    debts = {
        p["inn"]: _rub(_metrics(p).get(DEBT), str(_inputs(p).get("unit") or "")) for p in others
    }
    ranked = sorted(
        (p for p in others if debts[p["inn"]] is not None),
        key=lambda p: (-debts[p["inn"]], p["inn"]),
    )
    added = ranked[: max(SAMPLE_SIZE - len(pool), 0)]
    picked = [*pool, *added]
    place = {p["inn"]: number for number, p in enumerate(ranked, start=1)}
    _sheet(
        book,
        "Выборка",
        [
            (inn, [*values, "крупный долг по облигациям" if inn in systemic
                   else f"долг по отчётности, место {place[inn]}"])
            for inn, values in (line(point) for point in picked)
        ],
        cards_dir,
        ("основание отбора",),
    )
    carded = sum(1 for inn in points if (cards_dir / f"{inn}.md").exists())
    silent = sum(1 for inn in points if not any((bonds.get(inn) or {}).values()))
    disagree = sum(
        1
        for inn in points
        if not (bonds.get(inn) or {}).get(OWN_FILE)
        and any((bonds.get(inn) or {}).get(name) for name in (LISTED, BOARD, BACKING))
    )
    summary: list[tuple[str, object]] = [
        ("Дата маршрута", f"{as_of:%d.%m.%Y}"),
        ("Источник", "история маршрута (routing_day: последняя точка дня), только чтение"),
        ("Эмитентов в маршруте", len(points)),
        ("из них с карточкой рядом с пакетом", carded),
        ("«Разбор»", len(by_basket["review"])),
        ("«Без внимания»", len(clear)),
        (
            "Крупный долг",
            f"верхние {routing.systemic.top_share:.0%} эмитентов по объёму облигаций "
            f"в обращении (routing.yaml, systemic.top_share): {len(systemic)} "
            f"из {len(volumes)} с раскрытым объёмом",
        ),
        ("«Без внимания» с крупным долгом по облигациям", len(pool)),
        (
            "Добор по долгу по отчётности",
            f"{len(added)} с наибольшим совокупным долгом базы ({DEBT}, в рублях) "
            f"из {len(ranked)} «Без внимания» без крупного долга по облигациям; "
            f"долг не раскрыт либо единица неизвестна у {len(others) - len(ranked)} — "
            "в добор не идут",
        ),
        ("В выборке", len(picked)),
        (
            "Нехватка выборки",
            f"нужно {SAMPLE_SIZE}, есть {len(picked)}" if len(picked) < SAMPLE_SIZE else "нет",
        ),
        ("Отбор", "решение владельца 30.09.2026; отбор для проверки, не правило маршрута; "
                  "детерминирован — добор по убыванию долга, ничья по ИНН"),
        (
            "Выпуски в обращении по источникам",
            f"графа «файл / перечень / срез» — тем же кодом, что периметр "
            f"(routing_store.bond_sources); молчат все источники у {silent} "
            f"из {len(points)}; файл эмитента пуст, а другой источник видит "
            f"выпуски, — у {disagree}",
        ),
        ("Ограничение", "объём облигаций в обращении, наименования, файл выпусков "
                        "и перечень Cbonds — с диска на день сборки, а не на дату "
                        "маршрута: истории объёмов и статусов нет; срез биржи — "
                        "последний не позже даты маршрута"),
    ]
    sheet = book.create_sheet("Сводка", 0)
    for label, value in summary:
        sheet.append([label, value])
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 100
    _instruction(book, as_of, len(by_basket["review"]), len(clear), len(picked))
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(out))
    return summary


def main(argv: list[str] | None = None) -> int:
    """Пишет пакет на дату и печатает сводку; дня в истории нет либо файл есть — 1."""
    logging.basicConfig(level=logging.ERROR, format="%(message)s")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--as-of", type=date.fromisoformat, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args(argv)
    if args.out.exists():
        print(f"пакет {args.out} уже есть и не переписывается: выберите другой путь")
        return 1
    try:
        with connection() as conn:
            summary = build(conn, args.as_of, args.out)
            conn.rollback()
    except NoRouteError as missing:
        print(missing)
        return 1
    print(f"{args.out}")
    for label, value in summary:
        print(f"- {label}: {value}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
